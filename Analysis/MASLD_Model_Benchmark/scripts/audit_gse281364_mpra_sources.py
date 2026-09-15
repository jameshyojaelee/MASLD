#!/usr/bin/env python3
"""Audit the authoritative GSE281364 MPRA source and local replicate topology."""

from __future__ import annotations

import argparse
import csv
from hashlib import sha256
import json
from pathlib import Path
import re

from openpyxl import load_workbook


class GSE281364SourceError(ValueError):
    """Raised when an MPRA source or topology differs from its requirements."""


def _sha256_file(path: Path) -> str:
    value = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(16 * 1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def _read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def _workbook_census(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    return [
        {"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column}
        for sheet in workbook.worksheets
    ]


def audit(arguments: argparse.Namespace) -> dict[str, object]:
    paths = [
        arguments.series_soft,
        arguments.sample_soft,
        arguments.ena_report,
        arguments.nature_html,
        arguments.supplementary_tables,
        arguments.source_data,
        arguments.fastq_r1,
        arguments.fastq_r2,
        arguments.local_counts_root,
        arguments.local_analysis_root,
    ]
    if arguments.output.exists() or any(path.is_symlink() for path in paths):
        raise GSE281364SourceError("GSE281364 source audit request differs")

    series = arguments.series_soft.read_text(encoding="utf-8")
    sample = arguments.sample_soft.read_text(encoding="utf-8")
    ena_rows = _read_tsv(arguments.ena_report)
    nature = arguments.nature_html.read_text(encoding="utf-8")
    expected_fastq = {
        "SRR31280365_1.fastq.gz": {
            "md5": "e2a206d1a8be823073db430665f514e7",
            "bytes": 4_199_316_437,
        },
        "SRR31280365_2.fastq.gz": {
            "md5": "e14106cd95597681759eded6e66c4e02",
            "bytes": 3_462_222_758,
        },
    }
    if (
        "^SERIES = GSE281364" not in series
        or "!Series_status = Public on Mar 16 2026" not in series
        or "!Series_pubmed_id = 42230773" not in series
        or len(re.findall(r"^!Series_sample_id = GSM", series, flags=re.MULTILINE))
        != 33
        or len(
            re.findall(
                r"^!Series_supplementary_file = .*DNA_RNA_counts\.txt\.gz$",
                series,
                flags=re.MULTILINE,
            )
        )
        != 16
        or "^SAMPLE = GSM8619256" not in sample
        or "!Sample_title = variant and barcode mapping reads" not in sample
        or "!Sample_relation = SRA: https://www.ncbi.nlm.nih.gov/sra?term=SRX26659333"
        not in sample
        or "Read 1 sequences were trimmed" not in sample
        or "-l 20 -g TCACCATGGTGGCTTTACCAACAG" not in sample
        or "subsequent 20\u202fbp sequences was treated as barcode sequences"
        not in sample
    ):
        raise GSE281364SourceError("GEO series or plasmid-library sample differs")
    if len(ena_rows) != 1:
        raise GSE281364SourceError("ENA raw-run census differs")
    ena = ena_rows[0]
    if (
        ena.get("run_accession") != "SRR31280365"
        or ena.get("experiment_accession") != "SRX26659333"
        or ena.get("library_layout") != "PAIRED"
        or ena.get("read_count") != "71771807"
        or ena.get("base_count") != "21531542100"
        or ena.get("fastq_md5")
        != "e2a206d1a8be823073db430665f514e7;e14106cd95597681759eded6e66c4e02"
        or ena.get("fastq_bytes") != "4199316437;3462222758"
    ):
        raise GSE281364SourceError("ENA raw-run metadata differs")
    for path in (arguments.fastq_r1, arguments.fastq_r2):
        contract = expected_fastq.get(path.name)
        if contract is None or path.stat().st_size != contract["bytes"]:
            raise GSE281364SourceError("raw plasmid-library FASTQ differs")

    if (
        "10.1038/s41588-026-02617-8" not in nature
        or "2026 The Author(s), under exclusive licence to Springer Nature America, Inc."
        not in nature
        or '"open":false' not in nature
        or "41588_2026_2617_MOESM4_ESM.xlsx" not in nature
        or "41588_2026_2617_MOESM5_ESM.xlsx" not in nature
    ):
        raise GSE281364SourceError("final publication or rights metadata differs")
    tables = _workbook_census(arguments.supplementary_tables)
    source_data = _workbook_census(arguments.source_data)
    expected_table_sheets = [
        "Supplementary Table 1",
        "Supplementary Table 2",
        "Supplementary Table3",
        "Supplementary Table 4",
        "Supplementary Table 5",
        "Supplementary Table 6",
        "Supplementary Table 7",
        "Supplementary Table 8",
        "Supplementary Table 9",
        "Supplementary Table 10",
    ]
    if [item["sheet"] for item in tables] != expected_table_sheets:
        raise GSE281364SourceError("final supplementary-table census differs")
    if any(item["rows"] >= 5_369 for item in tables + source_data):
        raise GSE281364SourceError("complete variant table unexpectedly appeared")

    checksum_path = arguments.local_counts_root / "sha256sum.txt"
    expected_local: dict[str, str] = {}
    for line in checksum_path.read_text(encoding="utf-8").splitlines():
        checksum, source_path = line.split(maxsplit=1)
        source_path = source_path.strip()
        expected_local[Path(source_path).name] = checksum
    count_files = sorted(arguments.local_counts_root.glob("*DNA_RNA_counts.txt.gz"))
    if len(count_files) != 16 or set(expected_local) != {path.name for path in count_files}:
        raise GSE281364SourceError("local replicate-file census differs")
    manifests = _read_tsv(arguments.local_analysis_root / "counts/HepG2.sample_manifest.tsv")
    manifests += _read_tsv(arguments.local_analysis_root / "counts/LX2.sample_manifest.tsv")
    observed_topology = {
        (row["cell_line"], row["condition"], int(row["replicate"]))
        for row in manifests
    }
    expected_topology = {
        ("HepG2", condition, replicate)
        for condition in ("control", "PAOA")
        for replicate in range(1, 5)
    } | {
        ("LX2", condition, replicate)
        for condition in ("control", "TGFb")
        for replicate in range(1, 5)
    }
    if observed_topology != expected_topology or len(manifests) != 16:
        raise GSE281364SourceError("local experimental-replicate topology differs")
    for row in manifests:
        path = Path(row["path"])
        if (
            path.name not in expected_local
            or row["sha256"] != expected_local[path.name]
            or int(row["size"]) != path.stat().st_size
        ):
            raise GSE281364SourceError("local replicate manifest differs")
    gate = json.loads((arguments.local_analysis_root / "gate_verdict.json").read_text())
    if (
        not gate.get("source_reproduction_complete")
        or not gate.get("source_reproduction_pass")
        or not gate.get("mpranalyze_completion")
        or gate.get("official_source_calls_authoritative") is not True
        or gate.get("mpranalyze_source_authoritative") is not False
    ):
        raise GSE281364SourceError("local source-reproduction gate differs")

    receipt = {
        "schema_version": "masld-bench-gse281364-source-audit-v1",
        "status": "pass",
        "dataset_id": "gse281364",
        "geo_accession": "GSE281364",
        "geo_public_date": "2026-03-16",
        "bioproject": "PRJNA1183729",
        "final_doi": "10.1038/s41588-026-02617-8",
        "final_publication_date": "2026-06-02",
        "final_publication_rights": "exclusive_springer_nature_source_terms_apply",
        "raw_data_access": "public_GEO_SRA_source_terms_apply",
        "redistribution": "source_reference_only_no_raw_redistribution",
        "raw_plasmid_library": {
            "gsm": "GSM8619256",
            "experiment": "SRX26659333",
            "run": "SRR31280365",
            "paired_reads": 71_771_807,
            "fastq": {
                path.name: {
                    "bytes": path.stat().st_size,
                    "sha256": _sha256_file(path),
                }
                for path in (arguments.fastq_r1, arguments.fastq_r2)
            },
            "read_1_observed_oligo_bases": 107,
            "read_2_barcode_bases": 20,
        },
        "supplementary_tables": tables,
        "source_data": source_data,
        "complete_5369_variant_ref_alt_table_present": False,
        "synthesized_oligo_sequence_table_present": False,
        "processed_construct_sequence_present": False,
        "sequence_join_state": "join_unresolved_pending_raw_read_reconstruction",
        "sequence_inference_from_interval_prohibited": True,
        "local_count_files": [
            {
                "name": path.name,
                "bytes": path.stat().st_size,
                "sha256": expected_local[path.name],
            }
            for path in count_files
        ],
        "experimental_replicates": 16,
        "contexts": ["HepG2_control", "HepG2_PAOA", "LX2_control", "LX2_TGFb"],
        "replicates_per_context": 4,
        "biological_unit": "experimental_replicate",
        "pairing": "same_sample_different_aliquot_DNA_RNA",
        "donor_count": 0,
        "replicate_independence_for_model_evaluation": False,
        "outcome_role": "exposed_development_MPRA_only",
        "eQTL_or_ieQTL_supervision": False,
        "champion_or_sealed_evaluation": False,
        "post_head_exposure_state": "continual_seen",
    }
    arguments.output.write_text(
        json.dumps(receipt, sort_keys=True, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(receipt, sort_keys=True))
    return receipt


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--series-soft", type=Path, required=True)
    parser.add_argument("--sample-soft", type=Path, required=True)
    parser.add_argument("--ena-report", type=Path, required=True)
    parser.add_argument("--nature-html", type=Path, required=True)
    parser.add_argument("--supplementary-tables", type=Path, required=True)
    parser.add_argument("--source-data", type=Path, required=True)
    parser.add_argument("--fastq-r1", type=Path, required=True)
    parser.add_argument("--fastq-r2", type=Path, required=True)
    parser.add_argument("--local-counts-root", type=Path, required=True)
    parser.add_argument("--local-analysis-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    audit(parser.parse_args())


if __name__ == "__main__":
    main()
