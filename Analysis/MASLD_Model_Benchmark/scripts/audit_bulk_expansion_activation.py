#!/usr/bin/env python3
"""Freeze participant joins, folds, and GENCODE v49 crosswalks for one bulk cohort."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import gzip
from hashlib import sha256
import json
from pathlib import Path
import re

from openpyxl import load_workbook


SOURCE_SHA = "883f870e2a76f78c39d257c24f8e795d2e74f1bd32aba785251b52aa7cb8229e"
MATRIX_SHA = "3f934ab9c6ee730c6bd0887b10d963f8210f274c36dc0878cb814c85026dc038"
GTF_SHA = "73bbbbd6eb2f114d1536f7cbf2339a653e1edc889b0cbe78992e92560b5aaba9"
EXPECTED_N = {"GSE260666": 16, "GSE268273": 109, "GSE274114": 39}
ATTRIBUTE = re.compile(r'(?P<key>[A-Za-z_]+) "(?P<value>[^"]*)";')


class BulkActivationError(RuntimeError):
    """Raised when a participant, matrix, or reference join differs."""


def sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def characteristics(row: dict[str, str]) -> dict[str, str]:
    output: dict[str, str] = {}
    for value in json.loads(row["characteristics_json"]):
        if ": " not in value:
            raise BulkActivationError("characteristic lacks key-value separator")
        key, item = value.split(": ", 1)
        normalized = key.strip().lower()
        if normalized in output:
            raise BulkActivationError(f"duplicate characteristic: {normalized}")
        output[normalized] = item.strip()
    return output


def relation_accessions(relations_json: str) -> tuple[str, str]:
    relations = json.loads(relations_json)
    biosamples = [
        match.group(1)
        for value in relations
        if (match := re.search(r"/biosample/(SAMN[0-9]+)/?$", value)) is not None
    ]
    sra = [
        match.group(1)
        for value in relations
        if (match := re.search(r"[?&]term=(SRX[0-9]+)$", value)) is not None
    ]
    if len(biosamples) != 1 or len(sra) != 1:
        raise BulkActivationError("BioSample or SRA relation differs")
    return biosamples[0], sra[0]


def validate_gse274114_sample_id(value: str) -> None:
    if re.fullmatch(r"(?:FFPE[0-9]+|105925-001-[0-9]{3}|106039-001-[0-9]{3})", value) is None:
        raise BulkActivationError("GSE274114 source sample ID differs")


def parse_gtf(path: Path) -> tuple[dict[str, dict[str, str]], dict[str, list[str]]]:
    by_stable: dict[str, dict[str, str]] = {}
    by_symbol: dict[str, list[str]] = defaultdict(list)
    with gzip.open(path, "rt", encoding="utf-8", errors="strict") as handle:
        for line in handle:
            if line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("\t")
            if len(fields) != 9 or fields[2] != "gene":
                continue
            attrs = {match["key"]: match["value"] for match in ATTRIBUTE.finditer(fields[8])}
            gene_id = attrs.get("gene_id")
            symbol = attrs.get("gene_name")
            if not gene_id or not symbol:
                raise BulkActivationError("GENCODE gene lacks gene_id or gene_name")
            stable = gene_id.split(".", 1)[0]
            if stable in by_stable:
                raise BulkActivationError(f"duplicate stable GENCODE gene ID: {stable}")
            by_stable[stable] = {
                "gencode_v49_gene_id": gene_id,
                "gencode_v49_stable_id": stable,
                "gencode_v49_gene_name": symbol,
                "gencode_v49_gene_type": attrs.get("gene_type", ""),
            }
            by_symbol[symbol].append(stable)
    if not by_stable:
        raise BulkActivationError("GENCODE v49 GTF has no genes")
    return by_stable, dict(by_symbol)


def read_source_rows(source: Path, cohort: str) -> list[dict[str, str]]:
    with (source / "samples" / f"{cohort}.tsv").open(
        encoding="utf-8", newline=""
    ) as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    if len(rows) != EXPECTED_N[cohort]:
        raise BulkActivationError("source sample count differs")
    return rows


def cohort_participants(rows: list[dict[str, str]], cohort: str) -> list[dict[str, str]]:
    participants: list[dict[str, str]] = []
    for row in rows:
        chars = characteristics(row)
        if cohort == "GSE260666":
            match = re.search(r"\[([^]]+)\]$", row["title"])
            if match is None:
                raise BulkActivationError("GSE260666 title participant ID differs")
            participant = match.group(1)
            group = chars["disease state"].lower().replace(" ", "_")
            fibrosis = "not_reported"
            sex = "not_reported"
            source_material = "fresh_frozen_or_source_unspecified_liver"
        elif cohort == "GSE268273":
            descriptions = json.loads(row["description_json"])
            if len(descriptions) != 1 or not re.fullmatch(r"AR[0-9]+", descriptions[0]):
                raise BulkActivationError("GSE268273 participant description differs")
            participant = descriptions[0]
            group = chars["case/control"].lower()
            fibrosis = chars["fibrosis degree"].upper()
            sex = chars["sex"].lower()
            source_material = "fresh_frozen_liver_biopsy"
        else:
            participant = chars["sample id"]
            validate_gse274114_sample_id(participant)
            group = chars["group"].upper()
            fibrosis = "F0_to_F2_only_source_design"
            sex = "not_reported"
            source_material = chars["material"].upper()
        biosample, sra = relation_accessions(row["relations_json"])
        participants.append(
            {
                "participant_id": participant,
                "sample_accession": row["accession"],
                "biosample_accession": biosample,
                "sra_experiment": sra,
                "group": group,
                "fibrosis": fibrosis,
                "sex": sex,
                "source_material": source_material,
                "modality_status": "observed",
                "biological_unit": "participant_candidate" if cohort == "GSE274114" else "participant",
            }
        )
    keys = ("participant_id", "sample_accession", "biosample_accession", "sra_experiment")
    for key in keys:
        if len({row[key] for row in participants}) != len(participants):
            raise BulkActivationError(f"participant join is not unique: {key}")
    return participants


def assign_folds(rows: list[dict[str, str]]) -> dict[str, int]:
    by_stratum: dict[tuple[str, str], list[str]] = defaultdict(list)
    for row in rows:
        by_stratum[(row["group"], row["fibrosis"])].append(row["participant_id"])
    totals = [0] * 5
    group_totals: dict[str, list[int]] = defaultdict(lambda: [0] * 5)
    output: dict[str, int] = {}
    for (group, fibrosis), participants in sorted(by_stratum.items()):
        for participant in sorted(participants):
            fold = min(range(5), key=lambda value: (group_totals[group][value], totals[value], value))
            output[participant] = fold
            group_totals[group][fold] += 1
            totals[fold] += 1
    return output


def matrix_axis(matrix: Path, cohort: str) -> tuple[list[str], list[str], str]:
    if cohort == "GSE260666":
        path = matrix / "raw" / "GSE260nnn" / "GSE260666_raw_counts.txt.gz"
        with gzip.open(path, "rt", encoding="utf-8", errors="strict") as handle:
            header = handle.readline().rstrip("\n").split("\t")
            features = [line.split("\t", 1)[0] for line in handle]
        return header[1:], features, "raw_nonnegative_integer_gene_counts"
    if cohort == "GSE268273":
        path = matrix / "raw" / "GSE268nnn" / "GSE268273_EGN_RNAseq_Normalized_data.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        iterator = worksheet.iter_rows(values_only=True)
        header = list(next(iterator))
        samples = [str(value) for value in header[3:] if value is not None]
        features = [str(row[0]) for row in iterator if row[0] is not None]
        workbook.close()
        return samples, features, "source_global_limma_voom_transformed_expression"
    candidates = sorted((matrix / "raw" / "source").glob("*_quant.genes.sf.gz"))
    if len(candidates) != 39:
        raise BulkActivationError("GSE274114 quant file count differs")
    samples = [path.name.split("_", 1)[0] for path in candidates]
    with gzip.open(candidates[0], "rt", encoding="utf-8", errors="strict") as handle:
        if handle.readline().rstrip("\n").split("\t") != [
            "Name", "Length", "EffectiveLength", "TPM", "NumReads"
        ]:
            raise BulkActivationError("GSE274114 quant header differs")
        features = [line.split("\t", 1)[0] for line in handle]
    return samples, features, "per_sample_salmon_TPM_estimated_counts_effective_length"


def crosswalk(
    features: list[str],
    cohort: str,
    by_stable: dict[str, dict[str, str]],
    by_symbol: dict[str, list[str]],
) -> list[dict[str, str]]:
    output: list[dict[str, str]] = []
    for feature in features:
        stable = feature.split(".", 1)[0]
        if cohort == "GSE260666":
            matches = by_symbol.get(feature, [])
            state = "unique_symbol" if len(matches) == 1 else "ambiguous_symbol" if matches else "unmapped_symbol"
            stable = matches[0] if len(matches) == 1 else ""
        else:
            matches = [stable] if stable in by_stable else []
            state = "stable_id" if matches else "unmapped_stable_id"
        target = by_stable.get(stable, {})
        output.append(
            {
                "source_feature_id": feature,
                "source_stable_id": feature.split(".", 1)[0] if cohort != "GSE260666" else "",
                "mapping_state": state,
                "gencode_v49_gene_id": target.get("gencode_v49_gene_id", ""),
                "gencode_v49_stable_id": target.get("gencode_v49_stable_id", ""),
                "gencode_v49_gene_name": target.get("gencode_v49_gene_name", ""),
                "gencode_v49_gene_type": target.get("gencode_v49_gene_type", ""),
            }
        )
    return output


def write_tsv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("x", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=tuple(rows[0]), delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cohort", choices=sorted(EXPECTED_N), required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--source-artifacts-sha256", default=SOURCE_SHA)
    parser.add_argument("--matrix", type=Path, required=True)
    parser.add_argument("--matrix-artifacts-sha256", default=MATRIX_SHA)
    parser.add_argument("--gtf", type=Path, required=True)
    parser.add_argument("--gtf-sha256", default=GTF_SHA)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if sha256_file(args.source / "ARTIFACTS.json") != args.source_artifacts_sha256:
        raise BulkActivationError("source ARTIFACTS SHA-256 differs")
    if sha256_file(args.matrix / "ARTIFACTS.json") != args.matrix_artifacts_sha256:
        raise BulkActivationError("matrix ARTIFACTS SHA-256 differs")
    if sha256_file(args.gtf) != args.gtf_sha256:
        raise BulkActivationError("GENCODE v49 GTF SHA-256 differs")
    rows = cohort_participants(read_source_rows(args.source, args.cohort), args.cohort)
    samples, features, measurement = matrix_axis(args.matrix, args.cohort)
    expected_axis = {
        row["participant_id"] if args.cohort in {"GSE260666", "GSE268273"} else row["sample_accession"]
        for row in rows
    }
    if set(samples) != expected_axis:
        raise BulkActivationError("matrix sample axis does not match participant join")
    folds = assign_folds(rows)
    for row in rows:
        row["outer_fold"] = str(folds[row["participant_id"]])
    by_stable, by_symbol = parse_gtf(args.gtf)
    gene_rows = crosswalk(features, args.cohort, by_stable, by_symbol)
    args.output.mkdir(parents=True, exist_ok=False)
    write_tsv(args.output / "participant_join.tsv", rows)
    write_tsv(args.output / "gene_crosswalk.tsv", gene_rows)
    mapping_counts = Counter(row["mapping_state"] for row in gene_rows)
    training_blockers = [
        "cross_cohort_participant_and_expression_fingerprint_deduplication",
        "task_specific_activation_and_preprocessing_contract",
        "model_specific_derivative_weight_rights_review",
    ]
    if args.cohort == "GSE268273":
        training_blockers.insert(0, "reprocess_raw_SRA_inside_outer_folds_because_public_matrix_is_global_limma_voom")
    if args.cohort == "GSE274114":
        training_blockers.insert(0, "resolve_source_sample_id_to_unique_biological_donor")
    biological_unit = "participant_candidate" if args.cohort == "GSE274114" else "participant"
    audit = {
        "schema_version": "masld-bench-bulk-expansion-activation-audit-v1",
        "status": (
            "pass_sample_join_reference_qc_participant_identity_and_training_blocked"
            if args.cohort == "GSE274114"
            else "pass_join_reference_qc_training_blocked"
        ),
        "cohort": args.cohort,
        "biological_unit": biological_unit,
        "participants": len(rows),
        "group_counts": dict(sorted(Counter(row["group"] for row in rows).items())),
        "fibrosis_counts": dict(sorted(Counter(row["fibrosis"] for row in rows).items())),
        "outer_fold_counts": dict(sorted(Counter(folds.values()).items())),
        "matrix_sample_axis_exact": True,
        "source_features": len(features),
        "mapping_counts": dict(sorted(mapping_counts.items())),
        "measurement": measurement,
        "normalization_fit_inside_outer_training_fold": args.cohort != "GSE268273",
        "model_training_activated": False,
        "remaining_blockers": training_blockers,
    }
    rights = {
        "schema_version": "masld-bench-public-bulk-rights-v1",
        "access_tier": "public",
        "new_dua_or_controlled_access_required": False,
        "internal_nonclinical_research_use": "allowed_by_project_public_data_definition",
        "explicit_data_license_detected": False,
        "raw_or_processed_data_redistribution": "prohibited_by_project_policy",
        "released_weight_redistribution": "requires_model_specific_terms_and_legal_review",
        "clinical_use": False,
    }
    (args.output / "activation_audit.json").write_text(json.dumps(audit, indent=2, sort_keys=True) + "\n")
    (args.output / "rights_contract.json").write_text(json.dumps(rights, indent=2, sort_keys=True) + "\n")
    print(json.dumps(audit, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
