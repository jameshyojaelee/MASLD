#!/usr/bin/env python3
"""Acquire and profile processed matrices for three public bulk expansion cohorts."""

from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import csv
import gzip
from hashlib import sha256
import json
import math
from pathlib import Path
import urllib.parse
import urllib.request

from openpyxl import load_workbook


EXPECTED_SOURCE_ARTIFACTS_SHA256 = (
    "883f870e2a76f78c39d257c24f8e795d2e74f1bd32aba785251b52aa7cb8229e"
)
SERIES_URLS = {
    "GSE260666": [
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE260nnn/GSE260666/suppl/GSE260666_raw_counts.txt.gz"
    ],
    "GSE268273": [
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE268nnn/GSE268273/suppl/GSE268273_EGN_RNAseq_Normalized_data.xlsx",
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE268nnn/GSE268273/suppl/GSE268273_EGN_RNAseq_Processed_data.xlsx",
    ],
}
MAX_FILE_BYTES = 1024 * 1024 * 1024


class BulkMatrixAuditError(RuntimeError):
    """Raised when a processed bulk source violates the acquisition contract."""


def sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_samples(source: Path, accession: str) -> list[dict[str, str]]:
    with (source / "samples" / f"{accession}.tsv").open(
        encoding="utf-8", newline=""
    ) as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    return rows


def https_url(value: str) -> str:
    prefix = "ftp://ftp.ncbi.nlm.nih.gov/"
    if value.startswith(prefix):
        return "https://ftp.ncbi.nlm.nih.gov/" + value[len(prefix) :]
    if value.startswith("https://ftp.ncbi.nlm.nih.gov/"):
        return value
    raise BulkMatrixAuditError("processed matrix URL is not an NCBI source")


def download(url: str, root: Path) -> dict[str, object]:
    url = https_url(url)
    accession = next((value for value in url.split("/") if value.startswith("GSE")), "source")
    directory = root / accession
    directory.mkdir(exist_ok=True)
    target = directory / Path(urllib.parse.urlparse(url).path).name
    if target.exists():
        raise BulkMatrixAuditError(f"refusing to overwrite matrix source: {target}")
    request = urllib.request.Request(
        url, headers={"User-Agent": "masld-bench-bulk-matrix-audit/1.0"}
    )
    digest = sha256()
    byte_count = 0
    with urllib.request.urlopen(request, timeout=300) as response:
        declared = response.headers.get("Content-Length")
        with target.open("xb") as handle:
            while True:
                block = response.read(8 * 1024 * 1024)
                if not block:
                    break
                byte_count += len(block)
                if byte_count > MAX_FILE_BYTES:
                    raise BulkMatrixAuditError(f"matrix source exceeds size cap: {url}")
                digest.update(block)
                handle.write(block)
        if declared is not None and int(declared) != byte_count:
            raise BulkMatrixAuditError(f"Content-Length differs: {url}")
    return {
        "url": url,
        "local_file": str(target.relative_to(root.parent)),
        "bytes": byte_count,
        "sha256": digest.hexdigest(),
    }


def numeric_profile(values: list[float]) -> dict[str, object]:
    finite = [value for value in values if math.isfinite(value)]
    return {
        "numeric_cells": len(values),
        "finite_numeric_cells": len(finite),
        "nonfinite_numeric_cells": len(values) - len(finite),
        "minimum": min(finite) if finite else None,
        "maximum": max(finite) if finite else None,
        "negative_cells": sum(value < 0 for value in finite),
        "nonnegative_integer_cells": sum(
            value >= 0 and value.is_integer() for value in finite
        ),
    }


def profile_gse260666(path: Path) -> dict[str, object]:
    with gzip.open(path, "rt", encoding="utf-8", errors="strict") as handle:
        header = handle.readline().rstrip("\n").split("\t")
        if len(header) != 17:
            raise BulkMatrixAuditError("GSE260666 raw-count header width differs")
        rows = 0
        values: list[float] = []
        gene_digest = sha256()
        for line_number, line in enumerate(handle, start=2):
            fields = line.rstrip("\n").split("\t")
            if len(fields) != len(header):
                raise BulkMatrixAuditError(
                    f"GSE260666 matrix width differs at line {line_number}"
                )
            rows += 1
            gene_digest.update(fields[0].encode("utf-8") + b"\n")
            try:
                values.extend(float(value) for value in fields[1:])
            except ValueError as error:
                raise BulkMatrixAuditError("GSE260666 has a nonnumeric count") from error
    profile = numeric_profile(values)
    if profile["nonfinite_numeric_cells"] or profile["negative_cells"]:
        raise BulkMatrixAuditError("GSE260666 has invalid raw counts")
    return {
        "feature_rows": rows,
        "sample_columns": 16,
        "header": header,
        "gene_axis_sha256": gene_digest.hexdigest(),
        **profile,
        "declared_measurement": "StringTie_raw_sequence_counts",
        "integer_count_likelihood_candidate": profile["nonnegative_integer_cells"]
        == profile["numeric_cells"],
    }


def profile_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, object] = {}
    for worksheet in workbook.worksheets:
        numeric: list[float] = []
        strings = 0
        first_rows: list[list[object]] = []
        populated_rows = 0
        max_populated_columns = 0
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            materialized = list(row)
            if row_index <= 5:
                first_rows.append(materialized[:12])
            populated = [value for value in materialized if value is not None]
            if populated:
                populated_rows += 1
                max_populated_columns = max(max_populated_columns, len(materialized))
            for value in populated:
                if isinstance(value, bool):
                    strings += 1
                elif isinstance(value, (int, float)):
                    numeric.append(float(value))
                else:
                    strings += 1
        sheets[worksheet.title] = {
            "worksheet_max_row": worksheet.max_row,
            "worksheet_max_column": worksheet.max_column,
            "populated_rows": populated_rows,
            "max_populated_columns": max_populated_columns,
            "string_or_boolean_cells": strings,
            "first_rows_first_12_columns": first_rows,
            **numeric_profile(numeric),
        }
    workbook.close()
    if not sheets:
        raise BulkMatrixAuditError(f"workbook has no sheets: {path}")
    return {"sheets": sheets}


def profile_gse274114(paths: list[tuple[str, Path]]) -> dict[str, object]:
    expected_header = ["Name", "Length", "EffectiveLength", "TPM", "NumReads"]
    gene_axis: str | None = None
    file_profiles: list[dict[str, object]] = []
    for accession, path in paths:
        with gzip.open(path, "rt", encoding="utf-8", errors="strict") as handle:
            header = handle.readline().rstrip("\n").split("\t")
            if header != expected_header:
                raise BulkMatrixAuditError(f"GSE274114 Salmon header differs: {accession}")
            rows = 0
            digest = sha256()
            numeric: list[float] = []
            for line_number, line in enumerate(handle, start=2):
                fields = line.rstrip("\n").split("\t")
                if len(fields) != 5:
                    raise BulkMatrixAuditError(
                        f"GSE274114 width differs: {accession}:{line_number}"
                    )
                rows += 1
                digest.update(fields[0].encode("utf-8") + b"\n")
                try:
                    numeric.extend(float(value) for value in fields[1:])
                except ValueError as error:
                    raise BulkMatrixAuditError("GSE274114 has nonnumeric output") from error
        observed_axis = digest.hexdigest()
        if gene_axis is None:
            gene_axis = observed_axis
        elif gene_axis != observed_axis:
            raise BulkMatrixAuditError("GSE274114 gene order differs among samples")
        profile = numeric_profile(numeric)
        if profile["nonfinite_numeric_cells"] or profile["negative_cells"]:
            raise BulkMatrixAuditError("GSE274114 has invalid Salmon output")
        file_profiles.append(
            {"sample_accession": accession, "feature_rows": rows, **profile}
        )
    if len(paths) != 39:
        raise BulkMatrixAuditError("GSE274114 does not have 39 quant files")
    return {
        "sample_files": len(paths),
        "header": expected_header,
        "gene_axis_sha256": gene_axis,
        "feature_row_counts": sorted(
            {int(profile["feature_rows"]) for profile in file_profiles}
        ),
        "measurement_contract": {
            "TPM": "relative_abundance_continuous",
            "NumReads": "Salmon_estimated_counts_continuous",
            "effective_length": "required_for_count_scale_reconstruction",
            "rounding_forbidden": True,
        },
        "file_profiles": file_profiles,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument(
        "--source-artifacts-sha256", default=EXPECTED_SOURCE_ARTIFACTS_SHA256
    )
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if sha256_file(args.source / "ARTIFACTS.json") != args.source_artifacts_sha256:
        raise BulkMatrixAuditError("source ARTIFACTS SHA-256 differs")
    sample_rows = {
        accession: read_samples(args.source, accession)
        for accession in ("GSE260666", "GSE268273", "GSE274114")
    }
    urls = [url for values in SERIES_URLS.values() for url in values]
    gse274_rows: list[tuple[str, str]] = []
    for row in sample_rows["GSE274114"]:
        files = json.loads(row["supplementary_files_json"])
        if len(files) != 1 or not files[0].endswith("_quant.genes.sf.gz"):
            raise BulkMatrixAuditError("GSE274114 sample processed-file set differs")
        urls.append(files[0])
        gse274_rows.append((row["accession"], https_url(files[0])))
    if len(set(urls)) != 42:
        raise BulkMatrixAuditError("processed bulk URL census differs")

    args.output.mkdir(parents=True, exist_ok=False)
    raw = args.output / "raw"
    raw.mkdir()
    receipts: list[dict[str, object]] = []
    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = {executor.submit(download, url, raw): url for url in urls}
        for future in as_completed(futures):
            receipts.append(future.result())
    receipts.sort(key=lambda value: str(value["url"]))
    by_url = {
        str(receipt["url"]): args.output / str(receipt["local_file"])
        for receipt in receipts
    }
    inventory = {
        "schema_version": "masld-bench-bulk-expansion-matrix-audit-v1",
        "status": "pass",
        "cohorts": {
            "GSE260666": profile_gse260666(by_url[SERIES_URLS["GSE260666"][0]]),
            "GSE268273_normalized": profile_workbook(
                by_url[SERIES_URLS["GSE268273"][0]]
            ),
            "GSE268273_processed": {
                **profile_workbook(by_url[SERIES_URLS["GSE268273"][1]]),
                "model_input_allowed": False,
                "reason": "source_derived_group_contrast_and_covariate_adjusted_output",
            },
            "GSE274114": profile_gse274114(
                [(accession, by_url[url]) for accession, url in gse274_rows]
            ),
        },
        "candidate_record_is_active_dataset": False,
        "participant_join_authoritative": False,
        "normalization_fit_in_outer_training_folds": False,
        "model_training_activated": False,
        "next_step": "cohort_specific_axis_join_reference_rights_overlap_and_training_activation_audits",
    }
    (args.output / "matrix_inventory.json").write_text(
        json.dumps(inventory, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    (args.output / "download_receipt.json").write_text(
        json.dumps(receipts, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(json.dumps(inventory, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
