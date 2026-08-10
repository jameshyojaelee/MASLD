#!/usr/bin/env python3
"""Audit deposited schemas and biological units without projecting frozen outcomes."""

from __future__ import annotations

import csv
import gzip
import json
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from public_functional_common import (
    CANDIDATE_ROOT,
    PROJECT_ROOT,
    md5_file,
    read_tsv,
    require_sealed,
    sha256_file,
    write_tsv,
)


SAMPLE_FIELDS = [
    "dataset_id", "assay", "sample_id", "source_accession", "biological_unit_id",
    "technical_unit_id", "condition", "timepoint", "lineage", "genotype", "model",
    "treatment", "replicate", "include_in_inference", "exclusion_reason", "source_basis",
]
GATE_FIELDS = [
    "dataset_id", "gate_id", "status", "inference_authorized", "n_source_samples",
    "n_biological_units", "n_primary_complete_units", "source_limitation", "detail",
    "specification_sha256",
]


def parse_soft(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        for raw in handle:
            line = raw.rstrip("\n")
            if line.startswith("^SAMPLE = "):
                if current is not None:
                    records.append(current)
                current = {"accession": line.split(" = ", 1)[1], "characteristics": {}, "descriptions": []}
            elif current is not None and line.startswith("!Sample_title = "):
                current["title"] = line.split(" = ", 1)[1]
            elif current is not None and line.startswith("!Sample_description = "):
                current["descriptions"].append(line.split(" = ", 1)[1])
            elif current is not None and line.startswith("!Sample_characteristics_ch1 = "):
                value = line.split(" = ", 1)[1]
                key, separator, item = value.partition(":")
                if separator:
                    current["characteristics"][key.strip().lower()] = item.strip()
        if current is not None:
            records.append(current)
    return records


def counts_columns(path: Path) -> list[str]:
    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        return next(csv.reader(handle, delimiter="\t"))


def gate_pcls(root: Path, spec: str) -> tuple[list[dict[str, object]], dict[str, object], list[dict[str, object]]]:
    counts = root / "GSE200418_1047_uReads.txt.gz"
    soft = root / "GSE200418_family.soft.gz"
    audit: list[dict[str, object]] = []
    if not counts.is_file() or not soft.is_file():
        return [], gate("GSE200418", "PCLS_SOURCE", "skipped_source_unavailable", False, 0, 0, 0, "required file missing", "counts and SOFT are required", spec), audit
    records = parse_soft(soft)
    record_by_code: dict[str, dict[str, Any]] = {}
    for record in records:
        codes = [match.group(0) for value in record["descriptions"] for match in re.finditer(r"1047_\d{4}", value)]
        if len(set(codes)) == 1:
            record_by_code[codes[0]] = record
    headers = counts_columns(counts)
    sample_columns = []
    for column in headers:
        match = re.search(r"1047_\d{4}", column)
        if match:
            sample_columns.append((column, match.group(0)))
    rows: list[dict[str, object]] = []
    unmatched = []
    for column, code in sample_columns:
        record = record_by_code.get(code)
        if record is None:
            unmatched.append(code)
            continue
        characteristics = record["characteristics"]
        title = record.get("title", "")
        treatment_match = re.search(r"PCLS\s+(CTR|GFIPO|GFIP|GFIO|GFI|GF|G|F)\s+", title)
        treatment = treatment_match.group(1) if treatment_match else characteristics.get("treatment", "").rstrip("h")
        timepoint = characteristics.get("timepoint", "")
        donor = characteristics.get("donor id", "")
        rows.append(
            sample("GSE200418", "bulk_RNAseq_human_PCLS", column, record["accession"], donor, column, treatment, timepoint, "", "", "PCLS", treatment, "", True, "", "GEO SOFT donor/treatment/time plus deposited count-header join")
        )
    condition_pairs: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        if row["timepoint"] == "48h" and row["condition"] in {"GFIPO", "GFI"}:
            condition_pairs[str(row["biological_unit_id"])].add(str(row["condition"]))
    complete = sorted(donor for donor, values in condition_pairs.items() if values == {"GFIPO", "GFI"})
    passed = len(rows) == 99 and len({r["biological_unit_id"] for r in rows}) == 7 and len(unmatched) == 0 and len(complete) >= 4
    detail = f"mapped={len(rows)}/99; donors={len({r['biological_unit_id'] for r in rows})}; GFIPO-GFI 48h complete donors={len(complete)} ({','.join(complete)})"
    audit.extend([
        {"dataset_id": "GSE200418", "item": "expression_sample_columns", "value": len(sample_columns), "expected": 99, "pass": len(sample_columns) == 99},
        {"dataset_id": "GSE200418", "item": "soft_sample_records", "value": len(records), "expected": 99, "pass": len(records) == 99},
        {"dataset_id": "GSE200418", "item": "primary_complete_donors", "value": len(complete), "expected": ">=4", "pass": len(complete) >= 4},
    ])
    return rows, gate("GSE200418", "PCLS_SOURCE", "pass" if passed else "rejected_qc", passed, len(rows), len({r["biological_unit_id"] for r in rows}), len(complete), "", detail, spec), audit


def read_h5ad_vector(group: Any, name: str) -> list[str]:
    import h5py  # existing single-cell environment; no new dependency
    node = group[name]
    if isinstance(node, h5py.Dataset):
        values = node[()]
    elif isinstance(node, h5py.Group) and "codes" in node and "categories" in node:
        codes = node["codes"][()]
        categories = node["categories"][()]
        values = [categories[code] if code >= 0 else b"" for code in codes]
    else:
        raise TypeError(f"Unsupported H5AD obs encoding for {name}")
    output = []
    for value in values:
        if isinstance(value, bytes):
            output.append(value.decode("utf-8", errors="replace"))
        else:
            output.append(str(value))
    return output


def gate_schlo(root: Path, spec: str) -> tuple[list[dict[str, object]], dict[str, object], list[dict[str, object]]]:
    amended = (CANDIDATE_ROOT / "SOURCE_AMENDMENT_01.json").is_file() and (root / "GSE207889_os.h5ad").is_file()
    h5ad = root / ("GSE207889_os.h5ad" if amended else "GSE207889_masld-hlos_merged.h5ad")
    if not h5ad.is_file():
        return [], gate("GSE207889", "SCHLO_SOURCE", "skipped_source_unavailable", False, 0, 0, 0, "materialized H5AD missing", "compressed source must be downloaded and expanded", spec), []
    import h5py
    with h5py.File(h5ad, "r") as handle:
        obs = handle["obs"]
        obs_names = sorted(obs.keys())
        lower = {name.lower(): name for name in obs_names}
        sample_candidates = [name for name in obs_names if any(token in name.lower() for token in ["sample", "batch", "library", "orig.ident"])]
        treatment_candidates = [name for name in obs_names if any(token in name.lower() for token in ["treat", "condition"])]
        lineage_candidates = [name for name in obs_names if any(token in name.lower() for token in ["cell_type", "celltype", "annotation", "cluster"])]
        raw_present = "raw" in handle or "counts" in handle.get("layers", {})
        n_cells = int(obs.attrs.get("_index", "") != "") and len(read_h5ad_vector(obs, obs.attrs["_index"].decode() if isinstance(obs.attrs["_index"], bytes) else obs.attrs["_index"]))
        field_values: dict[str, list[str]] = {}
        for name in sorted(set(sample_candidates + treatment_candidates + lineage_candidates)):
            try:
                field_values[name] = read_h5ad_vector(obs, name)
            except Exception:
                continue
        deposited_sample_rows = []
        if all(name in obs for name in ["sample", "detailed_condition", "celltype_category"]):
            sample_values = read_h5ad_vector(obs, "sample")
            condition_values = read_h5ad_vector(obs, "detailed_condition")
            lineage_values = read_h5ad_vector(obs, "celltype_category")
            for sample_id in sorted(set(sample_values)):
                indices = [i for i, value in enumerate(sample_values) if value == sample_id]
                conditions = sorted({condition_values[i] for i in indices})
                lineages = sorted({lineage_values[i] for i in indices})
                deposited_sample_rows.append((sample_id, conditions, lineages, len(indices)))
    source_records = parse_soft(root / "GSE207889_family.soft.gz") if (root / "GSE207889_family.soft.gz").is_file() else []
    injury = [record for record in source_records if any(token in record.get("title", "") for token in ["Control-for-TGFB1", "TGFB1-", "Control-for-PA", "PA-500", "Control-for-OA", "OA-"])]
    condition_map = {"CTRL-TGFB1": ("CONTROL_TGFB", "CONTROL"), "TGFB1": ("TGFB", "TGFB"), "CTRL-PA": ("CONTROL_PA", "CONTROL"), "PA": ("PA", "PA"), "CTRL-OA": ("CONTROL_OA", "CONTROL"), "OA500": ("OA", "OA"), "CTRL": ("UNMATCHED_CTRL", "CONTROL")}
    pairs = Counter()
    rows = []
    condition_seen = Counter()
    for sample_id, conditions, lineages, n_sample_cells in deposited_sample_rows:
        if len(conditions) != 1 or conditions[0] not in condition_map:
            continue
        condition, treatment = condition_map[conditions[0]]
        condition_seen[condition] += 1
        replicate = str(condition_seen[condition])
        include = condition != "UNMATCHED_CTRL"
        if include:
            pairs[condition] += 1
        rows.append(sample("GSE207889", "scRNAseq_human_liver_organoid", sample_id, "OS_H5AD", f"W01_{condition}_{replicate}", sample_id, condition, "day25", ";".join(lineages), "WT", "OS_HLO", treatment, replicate, include, "" if include else "not_in_preregistered_matched_injury_contrasts", "deposited OS-HLO sample, detailed_condition, and celltype_category fields"))
    published_md5_ok = (not amended) or md5_file(root / "GSE207889_os.h5ad.gz") == "cff21f7580603e603297b4fbd56b8f75"
    passed = amended and published_md5_ok and len(injury) == 12 and all(pairs[key] == 2 for key in ["CONTROL_TGFB", "TGFB", "CONTROL_PA", "PA", "CONTROL_OA", "OA"]) and bool(sample_candidates) and bool(lineage_candidates) and raw_present and len(deposited_sample_rows) == 14
    limitation = "one W01 iPSC line; two source replicates per treatment are the biological-unit ceiling"
    detail = f"source_object={h5ad.name}; deposited samples={len(deposited_sample_rows)} (12 prespecified injury/control, 2 unmatched control); H5AD cells={n_cells}; obs fields={','.join(obs_names)}; raw_counts={raw_present}; published_md5_ok={published_md5_ok}"
    audit = [
        {"dataset_id": "GSE207889", "item": "obs_fields", "value": ";".join(obs_names), "expected": "sample/treatment/celltype/QC", "pass": bool(sample_candidates and lineage_candidates)},
        {"dataset_id": "GSE207889", "item": "raw_count_storage", "value": raw_present, "expected": True, "pass": raw_present},
        {"dataset_id": "GSE207889", "item": "injury_libraries", "value": len(injury), "expected": 12, "pass": len(injury) == 12},
        {"dataset_id": "GSE207889", "item": "deposited_OS_H5AD_samples", "value": len(deposited_sample_rows), "expected": 14, "pass": len(deposited_sample_rows) == 14},
        {"dataset_id": "GSE207889", "item": "published_OS_H5AD_md5", "value": published_md5_ok, "expected": True, "pass": published_md5_ok},
    ]
    return rows, gate("GSE207889", "SCHLO_SOURCE", "pass_with_source_limit" if passed else "rejected_qc", passed, len(rows), 2, 2, limitation, detail, spec), audit


def gate_acmsd(root: Path, spec: str) -> tuple[list[dict[str, object]], dict[str, object], list[dict[str, object]]]:
    workbook = root / "GSE253380_mPH_HLO_tmm_normalized_counts.xlsx"
    if not workbook.is_file():
        return [], gate("GSE253380", "ACMSD_SOURCE", "skipped_source_unavailable", False, 0, 0, 0, "workbook missing", "", spec), []
    from openpyxl import load_workbook
    book = load_workbook(workbook, read_only=True, data_only=True)
    candidates: set[str] = set()
    identity_strings: list[str] = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 120), values_only=True):
            for value in row:
                text = str(value or "")
                candidates.update(re.findall(r"(?:CC|TT)_[HS]_[1-6]", text))
                if re.search(r"iPSC|cell.?line|72\.3|CR00012|1383D6", text, re.I):
                    identity_strings.append(text)
    samples = sorted(candidates)
    line_key_present = any(re.search(r"(?:CC|TT)_[HS]_[1-6].*(?:72\.3|CR00012|1383D6)|(?:72\.3|CR00012|1383D6).*(?:CC|TT)_[HS]_[1-6]", text) for text in identity_strings)
    rows = []
    for sample_id in samples:
        genotype, model_code, replicate = sample_id.split("_")
        treatment = "DMSO" if int(replicate) <= 3 else "TLC065"
        model = "HLO" if model_code == "H" else "sHLO"
        biological = f"public_line_unknown_{genotype}_{replicate}"
        rows.append(sample("GSE253380", "bulk_RNAseq_human_liver_organoid", sample_id, "workbook_HLO_sheet", biological, sample_id, f"{genotype}_{model}_{treatment}", "endpoint", "", genotype, model, treatment, replicate, True, "", "deposited HLO matrix; treatment assignment follows frozen source design"))
    schema_pass = len(samples) == 24
    status = "pass" if schema_pass and line_key_present else ("pass_model_specific_no_line_key" if schema_pass else "rejected_qc")
    authorized = schema_pass
    limitation = "none" if line_key_present else "no authoritative public sample-to-iPSC-line key; model-specific only and cannot satisfy independent human-replication gate"
    detail = f"HLO sample columns={len(samples)}; public line key={line_key_present}; sheets={','.join(book.sheetnames)}"
    audit = [
        {"dataset_id": "GSE253380", "item": "HLO_sample_columns", "value": len(samples), "expected": 24, "pass": schema_pass},
        {"dataset_id": "GSE253380", "item": "authoritative_line_key", "value": line_key_present, "expected": True, "pass": line_key_present},
    ]
    return rows, gate("GSE253380", "ACMSD_SOURCE", status, authorized, len(samples), "unknown" if not line_key_present else 3, 3 if schema_pass else 0, limitation, detail, spec), audit


def parse_series_matrix_metadata(path: Path) -> dict[str, list[str]]:
    rows: dict[str, list[str]] = {}
    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if line.startswith("!series_matrix_table_begin"):
                break
            if line.startswith("!Sample_"):
                values = next(csv.reader([line.rstrip("\n")], delimiter="\t"))
                rows.setdefault(values[0], values[1:])
                if values[0] == "!Sample_characteristics_ch1":
                    rows.setdefault("!Sample_characteristics_ch1_all", []).append(values[1:])
    return rows


def gate_biopsy(root: Path, spec: str) -> tuple[list[dict[str, object]], dict[str, object], list[dict[str, object]]]:
    matrix = root / "GSE106737_series_matrix.txt.gz"
    if not matrix.is_file():
        return [], gate("GSE106737", "BIOPSY_SOURCE", "skipped_source_unavailable", False, 0, 0, 0, "series matrix missing", "", spec), []
    metadata = parse_series_matrix_metadata(matrix)
    titles = metadata.get("!Sample_title", [])
    accessions = metadata.get("!Sample_geo_accession", [])
    characteristic_rows = metadata.get("!Sample_characteristics_ch1_all", [])
    group_values = []
    for values in characteristic_rows:
        if values and all(str(value).startswith("patient group:") for value in values):
            group_values = [str(value).split(":", 1)[1].strip() for value in values]
            break
    rows = []
    for accession, title, group in zip(accessions, titles, group_values):
        pair = re.match(r"Pair-(\d+)-(?:female|male)-(baseline|followup)", title, re.I)
        if pair:
            participant, timepoint = pair.groups()
            condition = {"1": "RYGB_responder", "2": "lifestyle_responder", "3": "lifestyle_nonresponder"}[group]
            include = True
            reason = ""
        else:
            participant_match = re.match(r"patients_(\d+)_", title, re.I)
            participant = f"group4_{participant_match.group(1)}" if participant_match else ""
            timepoint = "baseline"
            condition = "baseline_only"
            include = False
            reason = "group4_baseline_only_excluded_from_reversibility"
        rows.append(sample("GSE106737", "microarray_human_liver_biopsy", accession, accession, participant, accession, condition, timepoint, "", "", "human_biopsy", "RYGB" if group == "1" else "lifestyle" if group in {"2", "3"} else "none", "", include, reason, "GEO title and deposited patient-group field"))
    counts = Counter(row["condition"] for row in rows)
    pair_counts = Counter()
    for condition in ["RYGB_responder", "lifestyle_responder", "lifestyle_nonresponder"]:
        by_participant: dict[str, set[str]] = defaultdict(set)
        for row in rows:
            if row["condition"] == condition:
                by_participant[str(row["biological_unit_id"])].add(str(row["timepoint"]))
        pair_counts[condition] = sum(times == {"baseline", "followup"} for times in by_participant.values())
    passed = len(rows) == 111 and pair_counts == Counter({"RYGB_responder": 21, "lifestyle_responder": 10, "lifestyle_nonresponder": 10}) and counts["baseline_only"] == 29
    detail = f"samples={len(rows)}; complete pairs RYGB/LSI-R/LSI-NR={pair_counts['RYGB_responder']}/{pair_counts['lifestyle_responder']}/{pair_counts['lifestyle_nonresponder']}; baseline-only={counts['baseline_only']}"
    audit = [
        {"dataset_id": "GSE106737", "item": "total_samples", "value": len(rows), "expected": 111, "pass": len(rows) == 111},
        {"dataset_id": "GSE106737", "item": "complete_pair_counts", "value": f"{pair_counts['RYGB_responder']};{pair_counts['lifestyle_responder']};{pair_counts['lifestyle_nonresponder']}", "expected": "21;10;10", "pass": passed},
        {"dataset_id": "GSE106737", "item": "baseline_only_excluded", "value": counts["baseline_only"], "expected": 29, "pass": counts["baseline_only"] == 29},
    ]
    return rows, gate("GSE106737", "BIOPSY_SOURCE", "pass" if passed else "rejected_qc", passed, len(rows), 70, 41, "", detail, spec), audit


def gate_clcc1(root: Path, spec: str) -> tuple[list[dict[str, object]], dict[str, object], list[dict[str, object]]]:
    workbook = root / "media-1.xlsx"
    if not workbook.is_file():
        return [], gate("CLCC1", "CLCC1_SOURCE", "skipped_source_unavailable", False, 0, 0, 0, "checksum-verified workbook unavailable", "", spec), []
    size_ok = workbook.stat().st_size == 4_215_858
    md5 = md5_file(workbook)
    checksum_ok = md5 == "fd3fcd19feec6572ff7df90c33f9d906"
    sheets: list[str] = []
    schema_detail = "workbook not opened because checksum failed"
    separate_universes = False
    if checksum_ok and size_ok:
        from openpyxl import load_workbook
        book = load_workbook(workbook, read_only=True, data_only=True)
        sheets = book.sheetnames
        normalized = " ".join(sheets).lower()
        separate_universes = any(token in normalized for token in ["genome", "primary", "screen"]) and any(token in normalized for token in ["857", "retest", "validation", "secondary"])
        schema_detail = f"sheets={';'.join(sheets)}"
    passed = size_ok and checksum_ok and separate_universes
    status = "pass" if passed else "skipped_source_unavailable" if not checksum_ok else "rejected_qc"
    limitation = "" if passed else "exact source workbook or separable genome-wide/retest universe not verified"
    audit = [
        {"dataset_id": "CLCC1", "item": "size_bytes", "value": workbook.stat().st_size, "expected": 4_215_858, "pass": size_ok},
        {"dataset_id": "CLCC1", "item": "md5", "value": md5, "expected": "fd3fcd19feec6572ff7df90c33f9d906", "pass": checksum_ok},
        {"dataset_id": "CLCC1", "item": "separate_screen_universes", "value": separate_universes, "expected": True, "pass": separate_universes},
    ]
    return [], gate("CLCC1", "CLCC1_SOURCE", status, passed, "", "gene", "", limitation, schema_detail, spec), audit


def sample(dataset_id: str, assay: str, sample_id: str, accession: str, biological: str, technical: str, condition: str, timepoint: str, lineage: str, genotype: str, model: str, treatment: str, replicate: str, include: bool, reason: str, basis: str) -> dict[str, object]:
    return {
        "dataset_id": dataset_id, "assay": assay, "sample_id": sample_id,
        "source_accession": accession, "biological_unit_id": biological,
        "technical_unit_id": technical, "condition": condition, "timepoint": timepoint,
        "lineage": lineage, "genotype": genotype, "model": model, "treatment": treatment,
        "replicate": replicate, "include_in_inference": str(include).lower(),
        "exclusion_reason": reason, "source_basis": basis,
    }


def gate(dataset_id: str, gate_id: str, status: str, authorized: bool, n_source: object, n_bio: object, n_primary: object, limitation: str, detail: str, spec: str) -> dict[str, object]:
    return {
        "dataset_id": dataset_id, "gate_id": gate_id, "status": status,
        "inference_authorized": str(authorized).lower(), "n_source_samples": n_source,
        "n_biological_units": n_bio, "n_primary_complete_units": n_primary,
        "source_limitation": limitation, "detail": detail, "specification_sha256": spec,
    }


def main() -> None:
    seal = require_sealed()
    if not (CANDIDATE_ROOT / "UNSEALED.json").is_file():
        raise RuntimeError("Source gates cannot run before acquisition writes UNSEALED.json")
    specification = str(seal["specification_sha256"])
    if (CANDIDATE_ROOT / "SOURCE_AMENDMENT_01.json").is_file() and (CANDIDATE_ROOT / "source_gate_status.tsv").is_file():
        for name in ["source_gate_status.tsv", "source_schema_audit.tsv"]:
            source = CANDIDATE_ROOT / name
            destination = CANDIDATE_ROOT / name.replace(".tsv", "_pre_source_amendment.tsv")
            if source.is_file() and not destination.exists():
                shutil.copyfile(source, destination)
    all_samples: list[dict[str, object]] = []
    gates: list[dict[str, object]] = []
    audits: list[dict[str, object]] = []
    functions = [
        ("GSE200418", gate_pcls), ("GSE207889", gate_schlo),
        ("GSE253380", gate_acmsd), ("GSE106737", gate_biopsy), ("CLCC1", gate_clcc1),
    ]
    for dataset_id, function in functions:
        rows, status, dataset_audits = function(CANDIDATE_ROOT / "sources" / dataset_id, specification)
        all_samples.extend(rows)
        gates.append(status)
        audits.extend(dataset_audits)
    gates.append(gate("MYOJIN_HLF", "MYOJIN_REUSE", "pass_reuse_only", True, 7762, "gene", 7762, "assay-specific survival non-support; no re-mining", "exact PHASE_C_VALIDATED snapshot frozen", specification))

    write_tsv(CANDIDATE_ROOT / "sample_manifest.tsv", all_samples, SAMPLE_FIELDS)
    write_tsv(CANDIDATE_ROOT / "source_gate_status.tsv", gates, GATE_FIELDS)
    write_tsv(CANDIDATE_ROOT / "source_schema_audit.tsv", audits, ["dataset_id", "item", "value", "expected", "pass"])
    assays = read_tsv(CANDIDATE_ROOT / "frozen_spec/public_assays.tsv")
    write_tsv(CANDIDATE_ROOT / "public_assay_registry.tsv", assays, list(assays[0]))
    source_manifests = []
    for path in sorted((CANDIDATE_ROOT / "sources").glob("*/source_manifest.tsv")):
        source_manifests.extend(read_tsv(path))
    if source_manifests:
        write_tsv(CANDIDATE_ROOT / "source_manifest.tsv", source_manifests, list(source_manifests[0]))
    print(json.dumps({"gates": gates, "n_sample_rows": len(all_samples)}, indent=2))


if __name__ == "__main__":
    main()
