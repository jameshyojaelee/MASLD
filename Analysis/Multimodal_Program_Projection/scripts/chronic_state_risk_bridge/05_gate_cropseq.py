#!/usr/bin/env python3
"""Gate GSE281160 guide joins and conservative risk-allele orientation."""

from __future__ import annotations

import csv
import gzip
import json
import math
import re
import time
import urllib.request
from collections import Counter
from pathlib import Path

import openpyxl

from bridge_common import (
    CANDIDATE_ROOT,
    atomic_write_text,
    read_tsv,
    require_validated_seal,
    sha256_file,
    write_tsv,
)


def finite(value: object) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def crop_source_tables() -> tuple[dict[str, dict[str, object]], dict[str, dict[str, object]], list[dict[str, str]]]:
    path = CANDIDATE_ROOT / "sources/ZHU_SUPP/Zhu2026_supplementary_tables.xlsx"
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    crop: dict[str, dict[str, object]] = {}
    for row in book["Supplementary Table 6"].iter_rows(min_row=5, values_only=True):
        if row[0]:
            crop[str(row[0])] = {
                "ctrl_mpra_log2fc": row[1], "ctrl_mpra_fdr": row[2],
                "paoa_mpra_log2fc": row[3], "paoa_mpra_fdr": row[4],
                "target_genes": str(row[9] or "").rstrip(";"),
            }
    registry: dict[str, dict[str, object]] = {}
    for row in book["Supplementary Table 7"].iter_rows(min_row=4, values_only=True):
        if row[0]:
            registry[str(row[0])] = {
                "chromosome": row[1], "position": row[2], "nea": row[4], "ea": row[5],
                "gwas_beta": row[6], "gwas_p": row[7], "table7_mpra_beta": row[8],
                "table7_mpra_p": row[9], "table7_mpra_fdr": row[10],
            }
    coordinates: dict[str, tuple[str, int]] = {}
    for sheet in ("Supplementary Table 2", "Supplementary Table 4"):
        for row in book[sheet].iter_rows(min_row=4, values_only=True):
            rsid = str(row[4] or "")
            if rsid not in crop:
                continue
            match = re.fullmatch(r"chr([^:]+):(\d+)-(\d+)", str(row[0]))
            if match:
                coordinates[rsid] = (match.group(1), (int(match.group(2)) + int(match.group(3))) // 2)
    guides: list[dict[str, str]] = []
    for row in book["Supplementary Table 9"].iter_rows(min_row=157, max_row=202, values_only=True):
        name, top = str(row[0] or ""), str(row[1] or "")
        if not name or not top.startswith("CACCG"):
            continue
        if name.startswith("sgi_rs"):
            match = re.fullmatch(r"sgi_(rs\d+)_(\d+)", name)
            if match:
                guides.append({"guide_name": name, "variant": match.group(1), "guide_number": match.group(2), "sequence": top[5:], "guide_type": "targeting"})
        elif name.startswith("sgihuCon_"):
            guides.append({"guide_name": name, "variant": "NTC", "guide_number": name.rsplit("_", 1)[1], "sequence": top[5:], "guide_type": "non_targeting"})
    for rsid, coordinate in coordinates.items():
        crop[rsid]["chromosome"], crop[rsid]["position"] = coordinate
    return crop, registry, guides


def dbsnp_alleles(rsid: str) -> tuple[str, set[str], str]:
    destination = CANDIDATE_ROOT / "sources/dbsnp" / f"{rsid}.json"
    destination.parent.mkdir(parents=True, exist_ok=True)
    if not destination.exists():
        request = urllib.request.Request(
            f"https://api.ncbi.nlm.nih.gov/variation/v0/refsnp/{rsid.removeprefix('rs')}",
            headers={"User-Agent": "MASLD-Plan43-risk-orientation/1.0"},
        )
        with urllib.request.urlopen(request, timeout=120) as response:
            payload = response.read()
        atomic_write_text(destination, payload.decode("utf-8"))
        time.sleep(0.4)
    data = json.loads(destination.read_text(encoding="utf-8"))
    placements = [item for item in data["primary_snapshot_data"]["placements_with_allele"] if item.get("is_ptlp")]
    if len(placements) != 1:
        raise RuntimeError(f"Expected one PTLP placement for {rsid}")
    spdis = [item["allele"]["spdi"] for item in placements[0]["alleles"]]
    references = {item["deleted_sequence"] for item in spdis}
    if len(references) != 1:
        raise RuntimeError(f"Nonunique reference allele for {rsid}")
    reference = next(iter(references))
    alternatives = {item["inserted_sequence"] for item in spdis if item["inserted_sequence"] != reference}
    return reference, alternatives, sha256_file(destination)


def guide_counts(guides: list[dict[str, str]]) -> tuple[list[dict[str, object]], dict[str, int]]:
    root = CANDIDATE_ROOT / "sources/GSE281160"
    with gzip.open(root / "GSE281160_barcodes.tsv.gz", "rt") as handle:
        matrix_barcodes = [line.strip() for line in handle if line.strip()]
    base_counts = Counter(value.rsplit("-", 1)[0] for value in matrix_barcodes)
    eligible_bases = {base for base, count in base_counts.items() if count == 1}
    assignments: dict[str, str] = {}
    with gzip.open(root / "GSE281160_cellAssign.txt.gz", "rt") as handle:
        for row in csv.DictReader(handle, delimiter="\t"):
            assignments[row["cell"].rsplit("-", 1)[0]] = row["barcode"]
    exact = {row["sequence"]: row["guide_name"] for row in guides if len(row["sequence"]) == 20}
    nineteen = {row["sequence"]: row["guide_name"] for row in guides if len(row["sequence"]) == 19}
    mapped: Counter[str] = Counter()
    unmapped: Counter[str] = Counter()
    for base in eligible_bases:
        sequence = assignments.get(base)
        if sequence is None:
            continue
        name = exact.get(sequence)
        reconciliation = "exact_20nt"
        if name is None:
            matches = [guide_name for prefix, guide_name in nineteen.items() if sequence.startswith(prefix)]
            if len(matches) == 1:
                name = matches[0]; reconciliation = "unique_19nt_supplement_prefix_to_20nt_assignment"
        if name is None:
            unmapped[sequence] += 1
        else:
            mapped[name] += 1
    rows: list[dict[str, object]] = []
    for guide in guides:
        count = mapped[guide["guide_name"]]
        rows.append({
            **guide, "n_matrix_cells": count, "minimum_50_pass": str(count >= 50).lower(),
            "sequence_reconciliation": "unique_19nt_supplement_prefix_to_20nt_assignment" if len(guide["sequence"]) == 19 else "exact_20nt",
        })
    summary = {
        "n_matrix_barcodes": len(matrix_barcodes),
        "n_unique_matrix_bases": len(eligible_bases),
        "n_ambiguous_matrix_bases": sum(count > 1 for count in base_counts.values()),
        "n_assignment_bases": len(assignments),
        "n_mapped_cells": sum(mapped.values()),
        "n_unregistered_guide_cells": sum(unmapped.values()),
    }
    return rows, summary


def main() -> None:
    seal = require_validated_seal()
    crop, registry, guides = crop_source_tables()
    loci = [row for row in read_tsv(CANDIDATE_ROOT / "corrected_genetic_locus_registry.tsv") if row["is_representative"] == "true"]
    guide_audit, guide_summary = guide_counts(guides)
    guide_by_variant: dict[str, list[dict[str, object]]] = {}
    for row in guide_audit:
        guide_by_variant.setdefault(str(row["variant"]), []).append(row)

    orientation_rows: list[dict[str, object]] = []
    dbsnp_manifest: list[dict[str, object]] = []
    for rsid in sorted(crop):
        source = crop[rsid]
        table7 = registry.get(rsid, {})
        beta = finite(table7.get("gwas_beta"))
        mpra = finite(source.get("ctrl_mpra_log2fc"))
        chrom = str(source.get("chromosome", ""))
        position = int(source["position"]) if source.get("position") is not None else None
        near = [] if position is None else [
            row for row in loci
            if row["chromosome"].removeprefix("chr") == chrom.removeprefix("chr")
            and abs(int(row["position"]) - position) <= 1_000_000
        ]
        targets = set(str(source["target_genes"]).split(";"))
        near.sort(key=lambda row: (row["representative_gene"] not in targets, abs(int(row["position"]) - (position or 0)), row["tier12_locus_uid"]))
        selected = near[0] if near else None
        reference = ""; alternatives: set[str] = set(); dbsnp_hash = ""; allele_match = False
        if beta is not None and table7.get("ea") and table7.get("nea"):
            try:
                reference, alternatives, dbsnp_hash = dbsnp_alleles(rsid)
                dbsnp_manifest.append({"rsid": rsid, "sha256": dbsnp_hash, "source": f"NCBI Variation API refsnp/{rsid.removeprefix('rs')}"})
                allele_match = {str(table7["ea"]), str(table7["nea"])} == ({reference} | alternatives)
            except Exception:
                allele_match = False
        risk_allele = str(table7.get("ea", "")) if beta is not None and beta > 0 else str(table7.get("nea", "")) if beta is not None and beta < 0 else ""
        risk_activity = ""
        if allele_match and mpra is not None and risk_allele:
            effect = mpra if risk_allele in alternatives else -mpra if risk_allele == reference else None
            if effect is not None:
                risk_activity = "increase" if effect > 0 else "decrease" if effect < 0 else "zero"
        source_guides = guide_by_variant.get(rsid, [])
        guide_identity_ok = len(source_guides) == 2
        guides_50 = guide_identity_ok and all(row["minimum_50_pass"] == "true" for row in source_guides)
        orientation_complete = bool(beta is not None and allele_match and risk_activity and selected and guide_identity_ok and guides_50)
        reasons = []
        if beta is None: reasons.append("no_finite_source_GWAS_beta")
        if not allele_match: reasons.append("allele_orientation_not_two_source_verified")
        if position is None: reasons.append("no_source_coordinate")
        if not selected: reasons.append("no_corrected_Tier12_locus_within_1Mb")
        if not guide_identity_ok: reasons.append("guide_variant_identity_mismatch_or_missing")
        if guide_identity_ok and not guides_50: reasons.append("one_or_more_guides_below_50_joined_cells")
        orientation_rows.append({
            "variant": rsid, "chromosome": chrom, "position_grch38": position or "",
            "source_target_genes": source["target_genes"], "nea": table7.get("nea", ""),
            "ea": table7.get("ea", ""), "gwas_beta": beta if beta is not None else "",
            "dbsnp_reference": reference, "dbsnp_alternatives": ";".join(sorted(alternatives)),
            "source_alleles_match_dbsnp": str(allele_match).lower(), "risk_allele": risk_allele,
            "ctrl_mpra_alt_minus_ref_log2fc": mpra if mpra is not None else "",
            "table6_table7_mpra_effect_agree": str(mpra is not None and finite(table7.get("table7_mpra_beta")) is not None and abs(mpra - float(table7["table7_mpra_beta"])) < 1e-8).lower(),
            "risk_allele_activity_direction": risk_activity,
            "nearest_corrected_locus_uid": selected["tier12_locus_uid"] if selected else "",
            "nearest_corrected_representative_gene": selected["representative_gene"] if selected else "",
            "distance_to_corrected_locus_bp": abs(int(selected["position"]) - position) if selected and position is not None else "",
            "n_source_guides": len(source_guides), "both_guides_50_cells": str(guides_50).lower(),
            "orientation_complete_for_bridge": str(orientation_complete).lower(),
            "exclusion_reasons": ";".join(reasons),
        })

    eligible = [row for row in orientation_rows if row["orientation_complete_for_bridge"] == "true"]
    eligible_loci = {row["nearest_corrected_locus_uid"] for row in eligible}
    directions = Counter(row["risk_allele_activity_direction"] for row in eligible)
    gate_pass = len(eligible_loci) >= 8 and directions["increase"] >= 3 and directions["decrease"] >= 3
    status = "pass" if gate_pass else "skipped_insufficient_oriented_loci"
    gate = [{
        "dataset_id": "GSE281160", "gate_id": "CROPSEQ_ORIENTATION",
        "status": status, "inference_authorized": str(gate_pass).lower(),
        "n_source_variants": len(crop), "n_complete_oriented_variants": len(eligible),
        "n_independent_corrected_tier12_loci": len(eligible_loci),
        "n_risk_activity_increase": directions["increase"],
        "n_risk_activity_decrease": directions["decrease"],
        "n_matrix_cells": guide_summary["n_matrix_barcodes"],
        "n_cells_with_registered_guide_join": guide_summary["n_mapped_cells"],
        "detail": "requires >=8 independent corrected Tier-1/2 loci, >=3 increasing and >=3 decreasing orientations, two registered guides, and >=50 joined cells per guide",
        "specification_sha256": seal["specification_sha256"],
    }]
    write_tsv(CANDIDATE_ROOT / "source_gates/cropseq_orientation_registry.tsv", orientation_rows, list(orientation_rows[0]))
    write_tsv(CANDIDATE_ROOT / "source_gates/cropseq_guide_audit.tsv", guide_audit, list(guide_audit[0]))
    write_tsv(CANDIDATE_ROOT / "source_gates/cropseq_dbsnp_manifest.tsv", dbsnp_manifest, list(dbsnp_manifest[0]) if dbsnp_manifest else ["rsid", "sha256", "source"])
    write_tsv(CANDIDATE_ROOT / "source_gates/cropseq_gate.tsv", gate, list(gate[0]))
    atomic_write_text(CANDIDATE_ROOT / "source_gates/CROPSEQ_TERMINAL_STATUS", status + "\n")
    print(json.dumps(gate[0], indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
