#!/usr/bin/env python3
"""Independently validate the fail-closed Hong GEN-03 public-source gate."""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import math
import subprocess
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Sequence

from genetics_common import (
    ContractError,
    DEFAULT_CANDIDATE_ROOT,
    PROJECT_ROOT,
    assert_candidate_root,
    atomic_write_tsv,
    clean,
    read_table,
    sha256_file,
    verify_no_quarantined_inputs,
    write_stage_seal,
)


PUBLIC_REL = Path("work/public_sources")
CONTEXT_REL = PUBLIC_REL / "zenodo_context"
EXPECTED_CONTEXTS = {
    "hepatocyte": {"disease_group3", "Hep-M3", "Hep-M4", "Hep-M6", "Hep-M8", "Hep-M12", "Hep-M14"},
    "cholangiocyte": {"disease_group3", "Chol-M1", "Chol-M2", "Chol-M3", "Chol-M4", "Chol-M6", "Chol-M7"},
    "stellate_cell": {"disease_group3", "HSC-M1", "HSC-M2", "HSC-M3"},
    "endothelial_cell": {"disease_group3", "endo-M1", "endo-M4", "endo-M7", "endo-M9"},
}
EXPECTED_RULES = {
    "bh_per_celltype_context",
    "bh_pooled_per_celltype",
    "bh_disease_separate_modules_pooled_per_celltype",
    "bh_pooled_global",
    "hierarchical_per_celltype_context_literal_r_boundary",
    "hierarchical_per_celltype_context_max_significant_best_q",
    "hierarchical_pooled_per_celltype_literal_r_boundary",
    "hierarchical_pooled_per_celltype_max_significant_best_q",
    "coefficient_bh_per_celltype_context",
    "coefficient_bh_pooled_per_celltype",
    "coefficient_bh_disease_separate_modules_pooled_per_celltype",
    "coefficient_bh_pooled_global",
    "coefficient_hierarchical_per_celltype_context_literal_r_boundary",
    "coefficient_hierarchical_per_celltype_context_max_significant_best_q",
    "coefficient_hierarchical_pooled_per_celltype_literal_r_boundary",
    "coefficient_hierarchical_pooled_per_celltype_max_significant_best_q",
}
EXPECTED_RELEASE = {
    "hong_covariate_availability.tsv",
    "hong_multiplicity_audit.tsv",
    "hong_producer_manifest.tsv",
    "hong_public_source_manifest.tsv",
    "hong_source_contract.tsv",
    "hong_source_gate_status.tsv",
    "hong_source_reproduction.tsv",
    "hong_tested_universe.tsv",
}
EXPECTED_PUBLIC_ASSETS = {
    "41588_2025_2237_MOESM1_ESM.pdf": (
        7_448_989,
        "2b16e3f0502e62a20baabd19b127afe8512eb51c2e0f2c865160e90a8a5d76ae",
    ),
    "41588_2025_2237_MOESM4_ESM.xlsx": (
        5_347_551,
        "2cfcfcbcdfa0deea0641595db04b8aa0c6ddb49eeaef8803b188821cb534cd91",
    ),
    "zenodo_record_14586466.json": (
        6_845,
        "d87c711f9cfd7213796bbe9ec52a92f6b8fd6e1232974ce57c0de3924c6b0345",
    ),
}
EXPECTED_GIT_COMMIT = "650e9fe46f57a0889b0b6ee0f7fd35f6fc7f4730"


def fail(condition: bool, message: str) -> None:
    if condition:
        raise ContractError(message)


def parse_probability(value: str, context: str) -> float:
    try:
        result = float(value)
    except ValueError as error:
        raise ContractError(f"non-numeric {context}: {value!r}") from error
    fail(not math.isfinite(result) or not 0 <= result <= 1, f"invalid {context}: {value!r}")
    return result


def parse_interaction_name(name: str) -> tuple[str, str]:
    prefix = "interaction."
    suffix = ".txt.gz"
    fail(not name.startswith(prefix) or not name.endswith(suffix), f"bad interaction name: {name}")
    stem = name[len(prefix) : -len(suffix)]
    for cell_type in sorted(EXPECTED_CONTEXTS, key=len, reverse=True):
        marker = f"{cell_type}_"
        if stem.startswith(marker):
            return cell_type, stem[len(marker) :]
    raise ContractError(f"unrecognized cell type in {name}")


def bh(values: Sequence[float]) -> list[float]:
    order = sorted(range(len(values)), key=lambda index: (values[index], index))
    result = [1.0] * len(values)
    running = 1.0
    for position in range(len(order) - 1, -1, -1):
        index = order[position]
        running = min(running, values[index] * len(values) / (position + 1), 1.0)
        result[index] = running
    return result


def flat_calls(
    records: Sequence[tuple[str, str, float, float]], statistic_index: int
) -> set[tuple[str, str]]:
    q_values = bh([record[statistic_index] for record in records])
    return {
        (record[0], record[1])
        for record, q_value in zip(records, q_values, strict=True)
        if q_value < 0.05
    }


def hierarchical_literal_calls(
    records: Sequence[tuple[str, str, float, float]], statistic_index: int
) -> set[tuple[str, str]]:
    grouped: dict[str, list[tuple[str, str, float, float]]] = defaultdict(list)
    for record in records:
        grouped[record[0]].append(record)
    genes = sorted(grouped)
    within: dict[str, list[tuple[tuple[str, str, float, float], float]]] = {}
    best = []
    for gene in genes:
        rows = grouped[gene]
        q_values = bh([record[statistic_index] for record in rows])
        paired = list(zip(rows, q_values, strict=True))
        within[gene] = paired
        best.append(min(q_value for _record, q_value in paired))
    gene_q = bh(best)
    significant = [index for index, q_value in enumerate(gene_q) if q_value < 0.05]
    if not significant:
        return set()
    maximum_gene_q = max(gene_q[index] for index in significant)
    boundary_index = next(
        index for index in significant if gene_q[index] == maximum_gene_q
    )
    boundary = best[boundary_index]
    return {
        (record[0], record[1])
        for gene in genes
        for record, q_value in within[gene]
        if q_value < boundary
    }


def source_union(candidate: Path) -> tuple[dict[str, set[tuple[str, str]]], dict[str, int]]:
    path = candidate / CONTEXT_REL / "significant_sc_eQTLs_with_annotations.txt.gz"
    flags = {cell_type: f"is_ieQTL_{cell_type}" for cell_type in EXPECTED_CONTEXTS}
    result = {cell_type: set() for cell_type in EXPECTED_CONTEXTS}
    rows = 0
    with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fail(reader.fieldnames is None, "source annotation has no header")
        for row in reader:
            rows += 1
            pair = clean(row["gene"]), clean(row["snp"])
            for cell_type, flag in flags.items():
                fail(row[flag] not in {"TRUE", "FALSE"}, f"non-boolean {flag}")
                if row[flag] == "TRUE":
                    result[cell_type].add(pair)
    counts = {
        "annotation_rows": rows,
        "source_union_flag_count": sum(len(value) for value in result.values()),
        "source_unique_pair_count": len(set().union(*result.values())),
        "source_iegenes": len({gene for value in result.values() for gene, _snp in value}),
    }
    fail(
        counts
        != {
            "annotation_rows": 21_828,
            "source_union_flag_count": 7_740,
            "source_unique_pair_count": 7_712,
            "source_iegenes": 2_136,
        },
        f"source annotation count drift: {counts}",
    )
    return result, counts


def read_interactions(
    candidate: Path,
) -> tuple[
    dict[tuple[str, str], list[tuple[str, str, float, float]]],
    list[dict[str, object]],
]:
    expected_table = {
        (row["cell_type"], row["context"]): row
        for row in read_table(candidate / "hong_tested_universe.tsv", "tsv")
    }
    fail(len(expected_table) != 23, "tested-universe table is not 23 unique contexts")
    results = {}
    checks = []
    observed_contexts: dict[str, set[str]] = defaultdict(set)
    for path in sorted((candidate / CONTEXT_REL).glob("interaction.*.txt.gz")):
        cell_type, context = parse_interaction_name(path.name)
        observed_contexts[cell_type].add(context)
        records = []
        terms: dict[tuple[str, str], Counter[str]] = defaultdict(Counter)
        lrt_by_pair: dict[tuple[str, str], set[str]] = defaultdict(set)
        coefficient_by_pair: dict[tuple[str, str], float] = {}
        phenotype = set()
        source_rows = 0
        with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            required = {"gene", "snp", "term", "pval_full", "lrt_pval", "Phenotype"}
            fail(reader.fieldnames is None or not required.issubset(reader.fieldnames), f"schema drift: {path.name}")
            for row in reader:
                source_rows += 1
                pair = clean(row["gene"]), clean(row["snp"])
                term = clean(row["term"])
                terms[pair][term] += 1
                lrt_by_pair[pair].add(clean(row["lrt_pval"]))
                phenotype.add(clean(row["Phenotype"]))
                if term.startswith("G:pheno"):
                    coefficient_by_pair[pair] = parse_probability(
                        row["pval_full"], f"{path.name} interaction coefficient p"
                    )
        fail(phenotype != {context}, f"phenotype mismatch: {path.name}")
        fail(any(len(values) != 1 for values in lrt_by_pair.values()), f"LRT mismatch: {path.name}")
        fail(
            any(counter["G"] != 1 or sum(term.startswith("G:pheno") for term in counter.elements()) != 1 for counter in terms.values()),
            f"coefficient term mismatch: {path.name}",
        )
        for gene, snp in terms:
            records.append(
                (
                    gene,
                    snp,
                    parse_probability(next(iter(lrt_by_pair[(gene, snp)])), f"{path.name} LRT p"),
                    coefficient_by_pair[(gene, snp)],
                )
            )
        expected = expected_table[(cell_type, context)]
        fail(len(records) != int(expected["tested_pairs"]), f"tested-pair drift: {path.name}")
        fail(source_rows != int(expected["source_rows"]), f"source-row drift: {path.name}")
        results[(cell_type, context)] = records
        checks.append(
            {
                "check_id": f"interaction_context_{cell_type}_{context}",
                "status": "pass",
                "observed": len(records),
                "expected": expected["tested_pairs"],
                "claim_authorized": "false",
                "note": "complete deposited interaction pair universe; cells remain nested in 44 donors",
            }
        )
    fail(dict(observed_contexts) != EXPECTED_CONTEXTS, f"context drift: {dict(observed_contexts)}")
    fail(sum(len(value) for value in results.values()) != 1_299_254, "state-specific tested-pair total drift")
    return results, checks


def normalize_variant(value: object) -> str:
    text = clean(value)
    pieces = text.split("_")
    if len(pieces) == 5 and pieces[-1].lower() == "hg38":
        return "_".join([pieces[0].removeprefix("chr"), *pieces[1:]])
    pieces = text.removeprefix("chr").split(":")
    if len(pieces) == 4:
        return "_".join(pieces) + "_hg38"
    if len(pieces) == 3 and len(pieces[2]) == 2:
        return f"{pieces[0]}_{pieces[1]}_{pieces[2][0]}_{pieces[2][1]}_hg38"
    raise ContractError(f"invalid source quartet variant: {text!r}")


def normalize_gene(value: object) -> str:
    if isinstance(value, datetime) and value.month == 12 and value.day == 1:
        return "DEC1"
    return clean(value)


def quartet_check(candidate: Path) -> dict[str, int]:
    from openpyxl import load_workbook

    path = candidate / CONTEXT_REL / "significant_quartets_n601.txt.gz"
    with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    zenodo = {
        (
            clean(row["celltype"]),
            clean(row["cell_state_module"]),
            clean(row["gene"]),
            normalize_variant(row["snp_chr_pos_ref_alt"]),
            clean(row["TF"]),
        )
        for row in rows
    }
    ids = {row["quartet_id"] for row in rows}
    workbook = load_workbook(
        candidate / PUBLIC_REL / "41588_2025_2237_MOESM4_ESM.xlsx",
        read_only=True,
        data_only=True,
    )
    sheet = workbook["SuppleTable 11"]
    workbook_rows = []
    workbook_ids = set()
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        workbook_ids.add(str(int(row[0])))
        workbook_rows.append(
            (
                clean(row[12]),
                clean(row[4]),
                normalize_gene(row[1]),
                normalize_variant(row[2]),
                clean(row[3]),
            )
        )
    workbook_set = set(workbook_rows)
    fail(
        (len(rows), len(ids), len(zenodo), len(workbook_rows), len(workbook_ids), len(workbook_set))
        != (640, 601, 601, 640, 601, 601),
        "quartet physical/unique count drift",
    )
    fail(zenodo != workbook_set, "Zenodo and Supplementary Table 11 quartet tuples differ")
    return {"physical_rows": 640, "unique_ids": 601, "unique_tuples": 601}


def donor_check(candidate: Path) -> dict[str, int]:
    from openpyxl import load_workbook

    workbook = load_workbook(
        candidate / PUBLIC_REL / "41588_2025_2237_MOESM4_ESM.xlsx",
        read_only=True,
        data_only=True,
    )
    sheet = workbook["SuppleTable 2"]
    rows = [
        row
        for row in sheet.iter_rows(min_row=3, values_only=True)
        if clean(row[0])
    ]
    included = [row for row in rows if row[16] is True]
    disease = Counter(clean(row[15]) for row in included)
    observed = {
        "source_donors_total": len(rows),
        "eqtl_donors_included": len(included),
        "eqtl_donors_excluded": len(rows) - len(included),
        "included_ctrl": disease["ctrl"],
        "included_masl": disease["MASL"],
        "included_early_mash": disease["eMASH"],
        "included_advanced_mash": disease["aMASH"],
    }
    expected = {
        "source_donors_total": 48,
        "eqtl_donors_included": 44,
        "eqtl_donors_excluded": 4,
        "included_ctrl": 23,
        "included_masl": 4,
        "included_early_mash": 7,
        "included_advanced_mash": 10,
    }
    fail(observed != expected, f"donor ascertainment drift: {observed}")
    reproduction = {
        row["metric"]: int(row["observed"])
        for row in read_table(candidate / "hong_source_reproduction.tsv", "tsv")
        if row["metric"] in expected
    }
    fail(reproduction != expected, f"reported donor reproduction drift: {reproduction}")
    return observed


def validate_manifests(candidate: Path) -> None:
    release = read_table(candidate / "hong_release_manifest.tsv", "tsv")
    fail({row["relative_path"] for row in release} != EXPECTED_RELEASE, "Hong release member drift")
    for row in release:
        path = candidate / row["relative_path"]
        fail(not path.is_file(), f"release member missing: {path}")
        fail(sha256_file(path) != row["sha256"], f"release hash drift: {path.name}")
        fail(path.stat().st_size != int(row["bytes"]), f"release byte drift: {path.name}")
        fail(row["canonical_promotion_status"] != "not_promoted", "canonical promotion asserted")
    seal_path = candidate / "work/stage_seals/06_hong_source_audit.json"
    seal = json.loads(seal_path.read_text(encoding="utf-8"))
    fail(seal.get("stage") != "06_hong_source_audit", "stage-06 seal identity drift")
    sealed = {row["relative_path"]: row for row in seal["outputs"]}
    expected_sealed = EXPECTED_RELEASE | {"hong_release_manifest.tsv"}
    fail(set(sealed) != expected_sealed, "stage-06 sealed output membership drift")
    for relative, row in sealed.items():
        path = candidate / relative
        fail(sha256_file(path) != row["sha256"], f"stage-06 hash drift: {relative}")
        fail(path.stat().st_size != int(row["bytes"]), f"stage-06 byte drift: {relative}")
    producers = read_table(candidate / "hong_producer_manifest.tsv", "tsv")
    fail(not producers, "Hong producer manifest is empty")
    for row in producers:
        path = PROJECT_ROOT / row["relative_path"]
        fail(not path.is_file(), f"producer missing: {row['relative_path']}")
        fail(sha256_file(path) != row["sha256"], f"producer hash drift: {row['relative_path']}")
        fail(path.stat().st_size != int(row["bytes"]), f"producer byte drift: {row['relative_path']}")


def validate_provenance_and_quarantine(candidate: Path) -> None:
    for name, (expected_bytes, expected_hash) in EXPECTED_PUBLIC_ASSETS.items():
        path = candidate / PUBLIC_REL / name
        fail(path.stat().st_size != expected_bytes, f"public asset byte drift: {name}")
        fail(sha256_file(path) != expected_hash, f"public asset hash drift: {name}")
    zenodo = json.loads(
        (candidate / PUBLIC_REL / "zenodo_record_14586466.json").read_text(
            encoding="utf-8"
        )
    )
    fail(str(zenodo.get("id")) != "14586466", "Zenodo record identity drift")
    fail(zenodo.get("doi") != "10.5281/zenodo.14586466", "Zenodo DOI drift")
    archive = zenodo.get("files", [{}])[0]
    fail(
        (
            archive.get("key"),
            archive.get("size"),
            archive.get("checksum"),
        )
        != (
            "Zenodo_250212.zip",
            6_647_735_352,
            "md5:1fd1a014829a8d6a0f2e97fcb07e69fb",
        ),
        "Zenodo archive contract drift",
    )
    code = candidate / PUBLIC_REL / "MASLD-sceQTL"
    commit = subprocess.run(
        ["git", "-C", str(code), "rev-parse", "HEAD"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    dirty = subprocess.run(
        ["git", "-C", str(code), "status", "--porcelain"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    fail(commit != EXPECTED_GIT_COMMIT or bool(dirty), "official code pin drift")
    inventory = read_table(candidate / "hong_public_archive_inventory.tsv", "tsv")
    fail(len(inventory) != 25, "bounded public inventory is not 25 entries")
    fail(Counter(row["role"] for row in inventory) != Counter({"interaction_summary": 23, "significant_call_annotation": 1, "quartet_table": 1}), "inventory role drift")
    paths = []
    for row in inventory:
        path = candidate / row["local_path"]
        paths.append(path)
        fail(row["retrieval_status"] != "crc_verified", f"CRC not verified: {path.name}")
        fail(row["published_archive_md5"] != "1fd1a014829a8d6a0f2e97fcb07e69fb", "archive MD5 provenance drift")
        fail(sha256_file(path) != row["local_sha256"], f"public-entry hash drift: {path.name}")
    sources = read_table(candidate / "hong_public_source_manifest.tsv", "tsv")
    fail(len(sources) != 29, "public source manifest is not 29 entries")
    for row in sources:
        if row["local_path"]:
            paths.append(candidate / row["local_path"])
    verify_no_quarantined_inputs(paths, PROJECT_ROOT)
    prohibited = {
        "context_rescue.tsv",
        "context_bridge.tsv",
        "context_evidence.tsv",
        "rescued_fraction.tsv",
        "gene_context_calls.tsv",
    }
    fail(any((candidate / name).exists() for name in prohibited), "claim-bearing GEN-03 output exists")
    manifest_text = "\n".join(
        (candidate / name).read_text(encoding="utf-8")
        for name in (
            "hong_public_archive_inventory.tsv",
            "hong_public_source_manifest.tsv",
            "hong_release_manifest.tsv",
        )
    )
    fail("ieqtl_disease_genes.csv" in manifest_text or "ieqtl_summary.csv" in manifest_text, "quarantined local outcome entered a Hong manifest")


def validate_multiplicity(
    candidate: Path,
    interactions: dict[tuple[str, str], list[tuple[str, str, float, float]]],
    union: dict[str, set[tuple[str, str]]],
) -> list[dict[str, object]]:
    table = {
        row["multiplicity_rule"]: row
        for row in read_table(candidate / "hong_multiplicity_audit.tsv", "tsv")
    }
    fail(set(table) != EXPECTED_RULES, "multiplicity-family membership drift")
    fail(any(row["reported_state_specific_call_count"] != "13648" for row in table.values()), "reported 13,648 denominator drift")
    fail(any(row["quartets_supported"] != "601" or row["quartets_expected"] != "601" for row in table.values()), "601-quartet support drift")
    fail(any(row["source_exact_reproduction"] == "true" for row in table.values()), "a multiplicity rule unexpectedly reproduced the source")

    checks = []
    rivals = (
        ("bh_per_celltype_context", flat_calls, 2),
        ("hierarchical_per_celltype_context_literal_r_boundary", hierarchical_literal_calls, 2),
        ("coefficient_hierarchical_per_celltype_context_literal_r_boundary", hierarchical_literal_calls, 3),
    )
    for rule, caller, statistic_index in rivals:
        state_count = 0
        called_union = {cell_type: set() for cell_type in EXPECTED_CONTEXTS}
        for (cell_type, _context), records in interactions.items():
            calls = caller(records, statistic_index)
            state_count += len(calls)
            called_union[cell_type].update(calls)
        symmetric_difference = sum(
            len(called_union[cell_type] ^ union[cell_type])
            for cell_type in EXPECTED_CONTEXTS
        )
        expected = table[rule]
        fail(state_count != int(expected["state_specific_call_count"]), f"independent call-count mismatch: {rule}")
        fail(symmetric_difference != int(expected["union_symmetric_difference"]), f"independent union mismatch: {rule}")
        fail(state_count == 13_648 and symmetric_difference == 0, f"rival unexpectedly reproduces source: {rule}")
        checks.append(
            {
                "check_id": f"multiplicity_rival_{rule}",
                "status": "pass",
                "observed": f"calls={state_count};union_symdiff={symmetric_difference}",
                "expected": "not_13648_or_not_exact_union",
                "claim_authorized": "false",
                "note": "independently rederived from full public interaction statistics",
            }
        )
    gates = {
        row["gate_id"]: row
        for row in read_table(candidate / "hong_source_gate_status.tsv", "tsv")
    }
    expected_gates = {
        "GEN03_PUBLIC_PROVENANCE": "ready",
        "GEN03_DONOR_CONTEXT_UNIVERSE": "ready",
        "GEN03_QUARTET_REPRODUCTION": "ready",
        "GEN03_MULTIPLICITY_REPRODUCTION": "coverage_limited",
        "GEN03_COMPLETE_CELLTYPE_EQTL_BACKGROUND": "coverage_limited",
        "GEN03_REQUIRED_MATCHING_COVARIATES": "not_ready",
        "GEN03_CONTEXT_RESCUE_AUTHORIZATION": "prohibited",
        "GEN03_NEGATIVE_CLAIM_AUTHORIZATION": "prohibited",
        "GEN03_OVERALL": "coverage_limited_terminal",
    }
    fail(
        {gate: gates.get(gate, {}).get("status") for gate in expected_gates}
        != expected_gates,
        "GEN-03 gate-state drift",
    )
    fail(any(row["claim_bearing"] != "false" for row in gates.values()), "claim-bearing gate emitted")
    return checks


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--candidate-root", type=Path, default=DEFAULT_CANDIDATE_ROOT)
    args = parser.parse_args()
    fail(args.project_root.resolve() != PROJECT_ROOT.resolve(), "project-root drift")
    candidate = assert_candidate_root(args.project_root, args.candidate_root)
    validate_manifests(candidate)
    validate_provenance_and_quarantine(candidate)
    union, union_counts = source_union(candidate)
    interactions, context_checks = read_interactions(candidate)
    quartet = quartet_check(candidate)
    donors = donor_check(candidate)
    multiplicity_checks = validate_multiplicity(candidate, interactions, union)
    report = [
        {
            "check_id": "public_provenance_and_quarantine",
            "status": "pass",
            "observed": "29_manifested_sources;25_crc_entries;0_quarantined_inputs",
            "expected": "29_manifested_sources;25_crc_entries;0_quarantined_inputs",
            "claim_authorized": "false",
            "note": "official Nature/Zenodo/GitHub sources only; no canonical promotion",
        },
        {
            "check_id": "source_union_annotation",
            "status": "pass",
            "observed": ";".join(f"{key}={value}" for key, value in union_counts.items()),
            "expected": "annotation_rows=21828;source_union_flag_count=7740;source_unique_pair_count=7712;source_iegenes=2136",
            "claim_authorized": "false",
            "note": "source annotation is descriptive only because multiplicity is not reconstructed",
        },
        {
            "check_id": "quartet_cross_source_identity",
            "status": "pass",
            "observed": f"{quartet['physical_rows']}_rows;{quartet['unique_tuples']}_unique_tuples",
            "expected": "640_rows;601_unique_tuples",
            "claim_authorized": "false",
            "note": "Zenodo and Supplementary Table 11 tuple sets are identical",
        },
        {
            "check_id": "donor_ascertainment",
            "status": "pass",
            "observed": ";".join(f"{key}={value}" for key, value in donors.items()),
            "expected": "48 total;44 included;23 ctrl;4 MASL;7 eMASH;10 aMASH",
            "claim_authorized": "false",
            "note": "all eQTL/ieQTL inference remains based on the 44 included biological donors",
        },
        *context_checks,
        *multiplicity_checks,
        {
            "check_id": "terminal_authorization",
            "status": "pass",
            "observed": "coverage_limited_terminal",
            "expected": "coverage_limited_terminal",
            "claim_authorized": "false",
            "note": "no rescued fraction, source-negative, context bridge, or biological gene call is authorized",
        },
    ]
    report_path = candidate / "hong_validation_report.tsv"
    atomic_write_tsv(report_path, report, list(report[0]))
    stage6 = candidate / "work/stage_seals/06_hong_source_audit.json"
    validation_manifest_path = candidate / "hong_validation_manifest.tsv"
    validation_manifest = [
        {
            "relative_path": str(report_path.relative_to(candidate)),
            "sha256": sha256_file(report_path),
            "bytes": report_path.stat().st_size,
            "role": "GEN03_independent_validation",
            "canonical_promotion_status": "not_promoted",
        },
        {
            "relative_path": str(stage6.relative_to(candidate)),
            "sha256": sha256_file(stage6),
            "bytes": stage6.stat().st_size,
            "role": "GEN03_validated_upstream_seal",
            "canonical_promotion_status": "not_promoted",
        },
    ]
    atomic_write_tsv(
        validation_manifest_path,
        validation_manifest,
        list(validation_manifest[0]),
    )
    write_stage_seal(
        candidate,
        "07_validate_hong_source_audit",
        [report_path, validation_manifest_path],
        [stage6],
    )
    print(
        "PASS: GEN-03 public source gate independently validated as "
        "coverage_limited_terminal"
    )


if __name__ == "__main__":
    main()
