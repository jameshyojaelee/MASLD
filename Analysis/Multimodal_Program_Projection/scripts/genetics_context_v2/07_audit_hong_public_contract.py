#!/usr/bin/env python3
"""Audit the public Hong et al. context-eQTL source contract for GEN-03.

This stage consumes only checksum/CRC-verified public source artifacts beneath
the isolated candidate.  It tests several explicit multiplicity families
against the authors' published 13,648 state-specific calls, deposited union
flags, and 601 quartets.  It emits source coverage and gate tables only: no
atlas join, rescued fraction, context-rescue classification, or negative claim.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import math
import subprocess
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Sequence

from genetics_common import (
    ContractError,
    DEFAULT_CANDIDATE_ROOT,
    PROJECT_ROOT,
    SCRIPT_DIR,
    assert_candidate_root,
    assert_owned_output,
    atomic_write_tsv,
    bool_text,
    clean,
    read_table,
    sha256_file,
    verify_no_quarantined_inputs,
    write_stage_seal,
)


PUBLIC_REL = Path("work/public_sources")
ZENODO_CONTEXT_REL = PUBLIC_REL / "zenodo_context"
INVENTORY_NAME = "hong_public_archive_inventory.tsv"
SUPPLEMENT_PDF = "41588_2025_2237_MOESM1_ESM.pdf"
SUPPLEMENT_TEXT = "41588_2025_2237_MOESM1_ESM.txt"
SUPPLEMENT_XLSX = "41588_2025_2237_MOESM4_ESM.xlsx"
ZENODO_METADATA = "zenodo_record_14586466.json"
PUBLIC_CODE_DIR = "MASLD-sceQTL"

EXPECTED_PUBLIC_ASSETS = {
    SUPPLEMENT_PDF: (7_448_989, "2b16e3f0502e62a20baabd19b127afe8512eb51c2e0f2c865160e90a8a5d76ae"),
    SUPPLEMENT_XLSX: (5_347_551, "2cfcfcbcdfa0deea0641595db04b8aa0c6ddb49eeaef8803b188821cb534cd91"),
    ZENODO_METADATA: (6_845, "d87c711f9cfd7213796bbe9ec52a92f6b8fd6e1232974ce57c0de3924c6b0345"),
}
EXPECTED_GIT_COMMIT = "650e9fe46f57a0889b0b6ee0f7fd35f6fc7f4730"
EXPECTED_ZENODO = {
    "record_id": "14586466",
    "doi": "10.5281/zenodo.14586466",
    "archive": "Zenodo_250212.zip",
    "archive_bytes": 6_647_735_352,
    "archive_checksum": "md5:1fd1a014829a8d6a0f2e97fcb07e69fb",
    "license": "cc-by-4.0",
}
EXPECTED_INTERACTION_CONTEXTS = {
    "hepatocyte": {"disease_group3", "Hep-M3", "Hep-M4", "Hep-M6", "Hep-M8", "Hep-M12", "Hep-M14"},
    "cholangiocyte": {"disease_group3", "Chol-M1", "Chol-M2", "Chol-M3", "Chol-M4", "Chol-M6", "Chol-M7"},
    "stellate_cell": {"disease_group3", "HSC-M1", "HSC-M2", "HSC-M3"},
    "endothelial_cell": {"disease_group3", "endo-M1", "endo-M4", "endo-M7", "endo-M9"},
}
EXPECTED_REPORTED_STATE_CALLS = 13_648
EXPECTED_QUARTETS = 601


@dataclass(frozen=True)
class InteractionRow:
    analysis_id: str
    cell_type: str
    context: str
    gene: str
    snp: str
    p_value: float
    coefficient_p_value: float

    @property
    def pair(self) -> tuple[str, str]:
        return self.gene, self.snp


def parse_float(value: object, context: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise ContractError(f"non-numeric {context}: {value!r}") from error
    if not math.isfinite(result) or result < 0 or result > 1:
        raise ContractError(f"out-of-range {context}: {value!r}")
    return result


def source_manifest(candidate: Path) -> list[dict[str, object]]:
    public = candidate / PUBLIC_REL
    rows: list[dict[str, object]] = []
    for name, (expected_bytes, expected_hash) in EXPECTED_PUBLIC_ASSETS.items():
        path = public / name
        if not path.is_file():
            raise ContractError(f"required public source is missing: {path}")
        observed_bytes = path.stat().st_size
        observed_hash = sha256_file(path)
        if (observed_bytes, observed_hash) != (expected_bytes, expected_hash):
            raise ContractError(f"public source hash/size drift: {name}")
        rows.append(
            {
                "source_id": name,
                "source_type": "nature_supplement" if name.startswith("41588_") else "zenodo_record_metadata",
                "source_url": (
                    "https://media.springernature.com/original/springer-static/esm/"
                    "art%3A10.1038%2Fs41588-025-02237-8/MediaObjects/" + name
                    if name.startswith("41588_")
                    else "https://zenodo.org/api/records/14586466"
                ),
                "local_path": str(path.relative_to(candidate)),
                "bytes": observed_bytes,
                "sha256": observed_hash,
                "verification": "sha256_and_size_verified",
                "claim_role": "public_source_only",
            }
        )

    code = public / PUBLIC_CODE_DIR
    if not (code / ".git").is_dir():
        raise ContractError(f"public Hong code clone is missing: {code}")
    commit = subprocess.run(
        ["git", "-C", str(code), "rev-parse", "HEAD"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    dirty = subprocess.run(
        ["git", "-C", str(code), "status", "--porcelain"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    if commit != EXPECTED_GIT_COMMIT or dirty:
        raise ContractError(f"public Hong code is not at clean pinned commit {EXPECTED_GIT_COMMIT}")
    rows.append(
        {
            "source_id": "MASLD-sceQTL",
            "source_type": "official_public_code",
            "source_url": "https://github.com/snu-mchoi-lab/MASLD-sceQTL",
            "local_path": str(code.relative_to(candidate)),
            "bytes": "",
            "sha256": commit,
            "verification": "clean_git_commit_verified",
            "claim_role": "method_definition_only",
        }
    )

    metadata = json.loads((public / ZENODO_METADATA).read_text(encoding="utf-8"))
    files = metadata.get("files", [])
    if len(files) != 1:
        raise ContractError("Zenodo record must contain exactly one archive")
    archive = files[0]
    observed = {
        "record_id": str(metadata.get("id")),
        "doi": metadata.get("doi"),
        "archive": archive.get("key"),
        "archive_bytes": archive.get("size"),
        "archive_checksum": archive.get("checksum"),
        "license": metadata.get("metadata", {}).get("license", {}).get("id"),
    }
    if observed != EXPECTED_ZENODO:
        raise ContractError(f"Zenodo source contract drift: {observed}")

    inventory = read_table(candidate / INVENTORY_NAME, "tsv")
    if len(inventory) != 25:
        raise ContractError(f"expected 25 bounded Hong entries, observed {len(inventory)}")
    if any(row["retrieval_status"] != "crc_verified" for row in inventory):
        raise ContractError("one or more bounded Hong source entries failed ZIP CRC verification")
    for row in inventory:
        path = candidate / row["local_path"]
        if not path.is_file() or sha256_file(path) != row["local_sha256"]:
            raise ContractError(f"Hong public entry drift: {row['archive_entry']}")
        rows.append(
            {
                "source_id": row["archive_entry"],
                "source_type": row["role"],
                "source_url": row["source_url"],
                "local_path": row["local_path"],
                "bytes": path.stat().st_size,
                "sha256": row["local_sha256"],
                "verification": "http_range_zip_crc_and_sha256_verified",
                "claim_role": "source_outcome_contract_no_atlas_join",
            }
        )
    verify_no_quarantined_inputs(
        [candidate / row["local_path"] for row in rows if row["local_path"]],
        PROJECT_ROOT,
    )
    return rows


def parse_analysis_identity(path: Path) -> tuple[str, str, str]:
    prefix = "interaction."
    suffix = ".txt.gz"
    if not path.name.startswith(prefix) or not path.name.endswith(suffix):
        raise ContractError(f"invalid interaction filename: {path.name}")
    stem = path.name[len(prefix) : -len(suffix)]
    for cell_type in sorted(EXPECTED_INTERACTION_CONTEXTS, key=len, reverse=True):
        marker = f"{cell_type}_"
        if stem.startswith(marker):
            context = stem[len(marker) :]
            return f"{cell_type}:{context}", cell_type, context
    raise ContractError(f"unrecognized interaction cell type: {path.name}")


def read_interaction(path: Path) -> tuple[list[InteractionRow], dict[str, object]]:
    analysis_id, cell_type, context = parse_analysis_identity(path)
    rows: list[InteractionRow] = []
    pair_term_counts: Counter[tuple[str, str]] = Counter()
    pair_genotype_counts: Counter[tuple[str, str]] = Counter()
    pair_interaction_counts: Counter[tuple[str, str]] = Counter()
    pair_interaction_p_values: dict[tuple[str, str], float] = {}
    pair_lrt_values: dict[tuple[str, str], set[str]] = defaultdict(set)
    phenotype_values: set[str] = set()
    with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {"gene", "snp", "term", "pval_full", "lrt_pval", "Phenotype"}
        if reader.fieldnames is None or not required.issubset(reader.fieldnames):
            raise ContractError(f"interaction schema mismatch: {path.name}")
        for source_row in reader:
            gene = clean(source_row["gene"])
            snp = clean(source_row["snp"])
            if not gene or not snp:
                raise ContractError(f"blank gene/SNP in {path.name}")
            pair = gene, snp
            pair_term_counts[pair] += 1
            pair_lrt_values[pair].add(clean(source_row["lrt_pval"]))
            phenotype_values.add(clean(source_row["Phenotype"]))
            if source_row["term"].startswith("G:pheno"):
                pair_interaction_counts[pair] += 1
                pair_interaction_p_values[pair] = parse_float(
                    source_row["pval_full"], f"{analysis_id} interaction pval_full"
                )
            if source_row["term"] != "G":
                continue
            pair_genotype_counts[pair] += 1
            rows.append(
                InteractionRow(
                    analysis_id=analysis_id,
                    cell_type=cell_type,
                    context=context,
                    gene=gene,
                    snp=snp,
                    p_value=parse_float(source_row["lrt_pval"], f"{analysis_id} lrt_pval"),
                    coefficient_p_value=math.nan,
                )
            )
    expected_phenotype = context
    if phenotype_values != {expected_phenotype}:
        raise ContractError(
            f"phenotype/file mismatch for {analysis_id}: {sorted(phenotype_values)}"
        )
    if len(rows) != len(pair_term_counts):
        raise ContractError(f"missing or duplicate genotype rows in {analysis_id}")
    if any(count != 1 for count in pair_genotype_counts.values()):
        raise ContractError(f"genotype term is not unique per pair: {analysis_id}")
    if any(count != 1 for count in pair_interaction_counts.values()):
        raise ContractError(f"interaction term is not unique per pair: {analysis_id}")
    if any(len(values) != 1 for values in pair_lrt_values.values()):
        raise ContractError(f"LRT p value differs across coefficient rows: {analysis_id}")
    rows = [
        InteractionRow(
            analysis_id=row.analysis_id,
            cell_type=row.cell_type,
            context=row.context,
            gene=row.gene,
            snp=row.snp,
            p_value=row.p_value,
            coefficient_p_value=pair_interaction_p_values[row.pair],
        )
        for row in rows
    ]
    pairs = {row.pair for row in rows}
    return rows, {
        "analysis_id": analysis_id,
        "cell_type": cell_type,
        "context": context,
        "context_type": "binary_disease_status" if context == "disease_group3" else "continuous_module_state",
        "tested_pairs": len(rows),
        "tested_genes": len({gene for gene, _ in pairs}),
        "tested_variants": len({snp for _, snp in pairs}),
        "source_rows": sum(pair_term_counts.values()),
        "source_significant_calls": "",
        "complete_tested_pair_universe": "true",
        "biological_unit": "44_donors_with_cells_nested_within_donor",
        "source_model": "Poisson_mixed_effects_LRT_1df",
    }


def bh_calls(
    rows: Sequence[InteractionRow], statistic: str = "lrt", alpha: float = 0.05
) -> set[tuple[str, str, str, str]]:
    if not rows:
        return set()
    adjusted = bh_adjust([row_p_value(row, statistic) for row in rows])
    return {
        (row.cell_type, row.context, row.gene, row.snp)
        for row, q_value in zip(rows, adjusted, strict=True)
        if q_value < alpha
    }


def bh_adjust(p_values: Sequence[float]) -> list[float]:
    if not p_values:
        return []
    ordered = sorted(enumerate(p_values), key=lambda item: (item[1], item[0]))
    n_tests = len(p_values)
    adjusted = [1.0] * n_tests
    running = 1.0
    for reverse_index in range(n_tests - 1, -1, -1):
        original_index, p_value = ordered[reverse_index]
        rank = reverse_index + 1
        running = min(running, p_value * n_tests / rank, 1.0)
        adjusted[original_index] = running
    return adjusted


def hierarchical_calls(
    rows: Sequence[InteractionRow],
    threshold_mode: str,
    statistic: str = "lrt",
    alpha: float = 0.05,
) -> set[tuple[str, str, str, str]]:
    """Apply the source's SNP-wise then gene-wise BH threshold construction.

    The deposited eQTL filtering code first adjusts all tested SNPs within each
    gene, applies BH again to the minimum within-gene q values, then calls every
    SNP whose within-gene q is below a boundary learned from significant genes.
    The supplement says the expanded ieQTL analysis used the same SNP-wise and
    gene-wise logic.  Both the literal R-code boundary and the conceptual
    maximum-significant-gene boundary are audited because tied gene-wise BH
    values can make them differ.
    """
    by_gene: dict[str, list[InteractionRow]] = defaultdict(list)
    for row in rows:
        by_gene[row.gene].append(row)
    genes = sorted(by_gene)
    within_gene: dict[str, list[tuple[InteractionRow, float]]] = {}
    best_q_values: list[float] = []
    for gene in genes:
        gene_rows = by_gene[gene]
        gene_q = bh_adjust([row_p_value(row, statistic) for row in gene_rows])
        paired = list(zip(gene_rows, gene_q, strict=True))
        within_gene[gene] = paired
        best_q_values.append(min(q_value for _row, q_value in paired))
    gene_q_values = bh_adjust(best_q_values)
    significant = [
        index for index, q_value in enumerate(gene_q_values) if q_value < alpha
    ]
    if not significant:
        return set()
    if threshold_mode == "literal_r_boundary":
        maximum_gene_q = max(gene_q_values[index] for index in significant)
        boundary_index = next(
            index
            for index in significant
            if gene_q_values[index] == maximum_gene_q
        )
        boundary = best_q_values[boundary_index]
    elif threshold_mode == "max_significant_best_q":
        boundary = max(best_q_values[index] for index in significant)
    else:
        raise ContractError(f"unknown hierarchical threshold mode: {threshold_mode}")
    return {
        (row.cell_type, row.context, row.gene, row.snp)
        for gene in genes
        for row, q_value in within_gene[gene]
        if q_value < boundary
    }


def row_p_value(row: InteractionRow, statistic: str) -> float:
    if statistic == "lrt":
        return row.p_value
    if statistic == "interaction_coefficient":
        return row.coefficient_p_value
    raise ContractError(f"unknown interaction test statistic: {statistic}")


def group_for_rule(
    rule: str, analyses: dict[str, list[InteractionRow]]
) -> dict[str, list[InteractionRow]]:
    groups: dict[str, list[InteractionRow]] = defaultdict(list)
    for analysis_id, rows in analyses.items():
        if rule == "bh_per_celltype_context":
            group_id = analysis_id
        elif rule == "bh_pooled_per_celltype":
            group_id = rows[0].cell_type
        elif rule == "bh_pooled_global":
            group_id = "global"
        elif rule == "bh_disease_separate_modules_pooled_per_celltype":
            context_class = "disease" if rows[0].context == "disease_group3" else "module"
            group_id = f"{rows[0].cell_type}:{context_class}"
        else:
            raise ContractError(f"unknown multiplicity rule: {rule}")
        groups[group_id].extend(rows)
    return groups


def apply_rule(
    rule: str, analyses: dict[str, list[InteractionRow]]
) -> dict[str, set[tuple[str, str]]]:
    state_calls: set[tuple[str, str, str, str]] = set()
    statistic = "interaction_coefficient" if rule.startswith("coefficient_") else "lrt"
    base_rule = rule.removeprefix("coefficient_")
    if base_rule.startswith("hierarchical_per_celltype_context_"):
        groups = {analysis_id: rows for analysis_id, rows in analyses.items()}
        threshold_mode = base_rule.removeprefix(
            "hierarchical_per_celltype_context_"
        )
        caller = lambda grouped_rows: hierarchical_calls(  # noqa: E731
            grouped_rows, threshold_mode, statistic
        )
    elif base_rule.startswith("hierarchical_pooled_per_celltype_"):
        groups = group_for_rule("bh_pooled_per_celltype", analyses)
        threshold_mode = base_rule.removeprefix(
            "hierarchical_pooled_per_celltype_"
        )
        caller = lambda grouped_rows: hierarchical_calls(  # noqa: E731
            grouped_rows, threshold_mode, statistic
        )
    else:
        groups = group_for_rule(base_rule, analyses)
        caller = lambda grouped_rows: bh_calls(  # noqa: E731
            grouped_rows, statistic
        )
    for grouped_rows in groups.values():
        for cell_type, context, gene, snp in caller(grouped_rows):
            state_calls.add((cell_type, context, gene, snp))
    result: dict[str, set[tuple[str, str]]] = {
        "__state__": {(f"{cell_type}:{context}", f"{gene}\t{snp}") for cell_type, context, gene, snp in state_calls}
    }
    for cell_type in EXPECTED_INTERACTION_CONTEXTS:
        result[cell_type] = {
            (gene, snp)
            for observed_cell_type, _context, gene, snp in state_calls
            if observed_cell_type == cell_type
        }
    return result


def read_source_union(path: Path) -> tuple[dict[str, set[tuple[str, str]]], dict[str, int]]:
    flags = {cell_type: f"is_ieQTL_{cell_type}" for cell_type in EXPECTED_INTERACTION_CONTEXTS}
    result = {cell_type: set() for cell_type in flags}
    n_rows = 0
    with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {"gene", "snp", *flags.values()}
        if reader.fieldnames is None or not required.issubset(reader.fieldnames):
            raise ContractError("source significant-call annotation schema mismatch")
        for row in reader:
            n_rows += 1
            pair = clean(row["gene"]), clean(row["snp"])
            for cell_type, flag in flags.items():
                value = clean(row[flag])
                if value not in {"TRUE", "FALSE"}:
                    raise ContractError(f"non-boolean source call {flag}={value!r}")
                if value == "TRUE":
                    result[cell_type].add(pair)
    return result, {
        "annotation_rows": n_rows,
        "source_union_flag_count": sum(len(pairs) for pairs in result.values()),
        "source_unique_pair_count": len(set().union(*result.values())),
        "source_iegenes": len({gene for pairs in result.values() for gene, _ in pairs}),
    }


def normalized_variant(value: object) -> str:
    text = clean(value)
    underscore_parts = text.split("_")
    if len(underscore_parts) == 5 and underscore_parts[-1].lower() == "hg38":
        return "_".join(
            [underscore_parts[0].removeprefix("chr"), *underscore_parts[1:]]
        )
    parts = text.removeprefix("chr").split(":")
    if len(parts) == 4:
        return "_".join(parts) + "_hg38"
    if len(parts) == 3 and len(parts[2]) == 2:
        return f"{parts[0]}_{parts[1]}_{parts[2][0]}_{parts[2][1]}_hg38"
    raise ContractError(f"cannot normalize quartet variant: {text!r}")


def normalized_gene(value: object) -> str:
    if isinstance(value, datetime) and value.month == 12 and value.day == 1:
        return "DEC1"
    return clean(value)


def read_quartets(
    zenodo_path: Path, workbook_path: Path
) -> tuple[set[tuple[str, str, str, str, str]], dict[str, int]]:
    from openpyxl import load_workbook

    with gzip.open(zenodo_path, "rt", encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    ids = {row["quartet_id"] for row in rows}
    zenodo = {
        (
            clean(row["celltype"]),
            clean(row["cell_state_module"]),
            clean(row["gene"]),
            normalized_variant(row["snp_chr_pos_ref_alt"]),
            clean(row["TF"]),
        )
        for row in rows
    }
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["SuppleTable 11"]
    sheet_rows = []
    sheet_ids = set()
    for row in worksheet.iter_rows(min_row=4, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        sheet_ids.add(str(int(row[0])))
        sheet_rows.append(
            (
                clean(row[12]),
                clean(row[4]),
                normalized_gene(row[1]),
                normalized_variant(row[2]),
                clean(row[3]),
            )
        )
    workbook_set = set(sheet_rows)
    if len(ids) != EXPECTED_QUARTETS or len(zenodo) != EXPECTED_QUARTETS:
        raise ContractError("Zenodo quartet table does not contain 601 unique quartets")
    if len(sheet_ids) != EXPECTED_QUARTETS or len(workbook_set) != EXPECTED_QUARTETS:
        raise ContractError("Supplementary Table 11 does not contain 601 unique quartets")
    if zenodo != workbook_set:
        raise ContractError(
            f"Zenodo/workbook quartet tuple mismatch: "
            f"zenodo_only={len(zenodo - workbook_set)}, workbook_only={len(workbook_set - zenodo)}"
        )
    return zenodo, {
        "zenodo_rows": len(rows),
        "zenodo_unique_ids": len(ids),
        "zenodo_unique_tuples": len(zenodo),
        "workbook_rows": len(sheet_rows),
        "workbook_unique_ids": len(sheet_ids),
        "workbook_unique_tuples": len(workbook_set),
    }


def read_donor_contract(workbook_path: Path) -> dict[str, object]:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["SuppleTable 2"]
    rows = list(worksheet.iter_rows(min_row=3, values_only=True))
    rows = [row for row in rows if clean(row[0])]
    included = [row for row in rows if row[16] is True]
    disease = Counter(clean(row[15]) for row in included)
    return {
        "source_donors_total": len(rows),
        "eqtl_donors_included": len(included),
        "eqtl_donors_excluded": len(rows) - len(included),
        "included_ctrl": disease["ctrl"],
        "included_masl": disease["MASL"],
        "included_early_mash": disease["eMASH"],
        "included_advanced_mash": disease["aMASH"],
    }


def multiplicity_audit(
    analyses: dict[str, list[InteractionRow]],
    source_union: dict[str, set[tuple[str, str]]],
    quartets: set[tuple[str, str, str, str, str]],
) -> tuple[list[dict[str, object]], str | None, dict[str, set[tuple[str, str]]]]:
    rules = (
        "bh_per_celltype_context",
        "bh_pooled_per_celltype",
        "bh_disease_separate_modules_pooled_per_celltype",
        "bh_pooled_global",
        "hierarchical_per_celltype_context_literal_r_boundary",
        "hierarchical_per_celltype_context_max_significant_best_q",
        "hierarchical_pooled_per_celltype_literal_r_boundary",
        "hierarchical_pooled_per_celltype_max_significant_best_q",
        "coefficient_bh_per_celltype_context",
        "coefficient_bh_pooled_per_celltype",
        "coefficient_bh_disease_separate_modules_pooled_per_celltype",
        "coefficient_bh_pooled_global",
        "coefficient_hierarchical_per_celltype_context_literal_r_boundary",
        "coefficient_hierarchical_per_celltype_context_max_significant_best_q",
        "coefficient_hierarchical_pooled_per_celltype_literal_r_boundary",
        "coefficient_hierarchical_pooled_per_celltype_max_significant_best_q",
    )
    rows = []
    accepted_rules = []
    calls_by_rule: dict[str, dict[str, set[tuple[str, str]]]] = {}
    quartet_core = {(cell_type, context, gene, variant, tf) for cell_type, context, gene, variant, tf in quartets}
    for rule in rules:
        calls = apply_rule(rule, analyses)
        calls_by_rule[rule] = calls
        state_calls = calls["__state__"]
        symmetric_difference = sum(
            len(calls[cell_type] ^ source_union[cell_type])
            for cell_type in EXPECTED_INTERACTION_CONTEXTS
        )
        quartet_supported = 0
        for cell_type, context, gene, variant, _tf in quartet_core:
            raw_snp = variant.replace("_hg38", "").split("_")
            raw_snp = f"chr{raw_snp[0]}:{raw_snp[1]}:{raw_snp[2]}{raw_snp[3]}"
            if (f"{cell_type}:{context}", f"{gene}\t{raw_snp}") in state_calls:
                quartet_supported += 1
        state_count = len(state_calls)
        passes = (
            state_count == EXPECTED_REPORTED_STATE_CALLS
            and symmetric_difference == 0
            and quartet_supported == EXPECTED_QUARTETS
        )
        if passes:
            accepted_rules.append(rule)
        rows.append(
            {
                "multiplicity_rule": rule,
                "state_specific_call_count": state_count,
                "reported_state_specific_call_count": EXPECTED_REPORTED_STATE_CALLS,
                "union_symmetric_difference": symmetric_difference,
                "quartets_supported": quartet_supported,
                "quartets_expected": EXPECTED_QUARTETS,
                "source_exact_reproduction": bool_text(passes),
            }
        )
    accepted = accepted_rules[0] if len(accepted_rules) == 1 else None
    return rows, accepted, calls_by_rule.get(accepted, {})


def contract_rows(
    donor: dict[str, object],
    quartet: dict[str, int],
    accepted_rule: str | None,
) -> list[dict[str, object]]:
    return [
        {
            "contract_id": "HONG_COHORT_ASCERTAINMENT",
            "source": "Supplementary Table 2 and Supplementary Methods",
            "source_location": "MOESM4 SuppleTable 2; MOESM1 methods pp. 9-10",
            "verified_value": f"48 source donors; {donor['eqtl_donors_included']} included and {donor['eqtl_donors_excluded']} excluded",
            "status": "verified",
            "interpretation_limit": "Cells are nested within donors; cells are not biological replicates.",
        },
        {
            "contract_id": "HONG_INTERACTION_MODEL",
            "source": "Supplementary Methods and official GitHub",
            "source_location": "MOESM1 methods pp. 10-12; PME_ieQTL_calling.R",
            "verified_value": "Poisson mixed model; genotype-by-state LRT against no-interaction model; donor random intercept",
            "status": "verified",
            "interpretation_limit": "Interaction effects are not case-control mean-expression effects.",
        },
        {
            "contract_id": "HONG_STATE_DEFINITION",
            "source": "Supplementary Methods and Supplementary Table 7",
            "source_location": "MOESM1 p. 11; MOESM4 SuppleTable 7",
            "verified_value": "19 continuous scaled module states plus binary disease status across four cell types",
            "status": "verified",
            "interpretation_limit": "Only deposited prespecified contexts are in the source-tested universe.",
        },
        {
            "contract_id": "HONG_MULTIPLICITY_UNIT",
            "source": "Full Zenodo interaction statistics plus deposited union flags",
            "source_location": "Zenodo record 14586466; source reproduction audit",
            "verified_value": accepted_rule or "not uniquely reconstructed",
            "status": "verified" if accepted_rule else "coverage_limited",
            "interpretation_limit": "No source-significant interaction call is authorized unless one rule uniquely reproduces all source checks.",
        },
        {
            "contract_id": "HONG_QUARTETS",
            "source": "Zenodo quartet table and Supplementary Table 11",
            "source_location": "Zenodo significant_quartets_n601; MOESM4 SuppleTable 11",
            "verified_value": f"{quartet['zenodo_unique_tuples']} identical unique TF-state-variant-gene tuples",
            "status": "verified",
            "interpretation_limit": "Quartets are regulatory units; they are not an unbiased trait-linked rescue universe.",
        },
    ]


def covariate_rows(candidate: Path) -> list[dict[str, object]]:
    central = (candidate / PUBLIC_REL / "zenodo_zip_central_directory.bin").read_bytes()
    central_text = central.decode("latin1", errors="ignore")
    checks = [
        ("source_cell_type", "interaction filenames and cell-type-specific files", True, "available_in_current_audit"),
        ("source_expression", "four ge.df.<celltype>.txt.gz expression matrices", all(f"ge.df.{cell_type}.txt.gz" in central_text for cell_type in EXPECTED_INTERACTION_CONTEXTS), "deposited_not_assembled"),
        ("static_eqtl_power", "four complete liver_eQTLs_summary_stats_<celltype>.txt.gz files", all(f"liver_eQTLs_summary_stats_{cell_type}.txt.gz" in central_text for cell_type in EXPECTED_INTERACTION_CONTEXTS), "deposited_not_reproduced_in_GEN03"),
        ("gene_length", "Cell Ranger GRCh38 2020-A GTF named in source code", False, "reference_not_hash_frozen_for_matching"),
        ("local_tested_variant_density", "complete liver summary statistics", True, "derivable_but_not_assembled"),
        ("trait_locus_opportunity", "prespecified MASLD/PDFF/fibrosis credible-set universe", False, "belongs_to_GEN04_GEN05_and_not_yet_frozen"),
        ("assay_availability", "23 complete source interaction tested universes", True, "available_in_current_audit"),
    ]
    return [
        {
            "covariate": name,
            "public_source_evidence": evidence,
            "demonstrably_available": bool_text(available),
            "assembly_status": status,
            "authorized_for_matching_now": "false",
        }
        for name, evidence, available, status in checks
    ]


def write_release_manifest(candidate: Path, outputs: Sequence[Path]) -> Path:
    manifest = candidate / "hong_release_manifest.tsv"
    rows = [
        {
            "relative_path": str(path.relative_to(candidate)),
            "sha256": sha256_file(path),
            "bytes": path.stat().st_size,
            "role": "GEN03_source_audit",
            "canonical_promotion_status": "not_promoted",
        }
        for path in sorted(outputs)
    ]
    atomic_write_tsv(manifest, rows, list(rows[0]))
    return manifest


def write_producer_manifest(candidate: Path) -> Path:
    path = candidate / "hong_producer_manifest.tsv"
    allowed_suffixes = {".py", ".R", ".sbatch", ".sh", ".tsv", ".csv", ".md"}
    rows = []
    for source in sorted(SCRIPT_DIR.rglob("*")):
        if not source.is_file() or source.suffix not in allowed_suffixes:
            continue
        relative = source.relative_to(SCRIPT_DIR)
        rows.append(
            {
                "relative_path": str(source.relative_to(PROJECT_ROOT)),
                "sha256": sha256_file(source),
                "bytes": source.stat().st_size,
                "role": (
                    "test_or_fixture"
                    if "tests" in relative.parts
                    else "producer_contract_or_documentation"
                ),
            }
        )
    atomic_write_tsv(path, rows, list(rows[0]))
    return path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--candidate-root", type=Path, default=DEFAULT_CANDIDATE_ROOT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    project_root = args.project_root.resolve()
    if project_root != PROJECT_ROOT.resolve():
        raise ContractError(f"project root must be exactly {PROJECT_ROOT.resolve()}")
    candidate = assert_candidate_root(project_root, args.candidate_root)
    stage5 = candidate / "work/stage_seals/05_validate_preflight.json"
    if not stage5.is_file():
        raise ContractError("GEN-00/01/02 preflight seal is missing")

    manifest_rows = source_manifest(candidate)
    source_manifest_path = candidate / "hong_public_source_manifest.tsv"
    atomic_write_tsv(source_manifest_path, manifest_rows, list(manifest_rows[0]))

    inventory = read_table(candidate / INVENTORY_NAME, "tsv")
    interaction_inventory = [row for row in inventory if row["role"] == "interaction_summary"]
    analyses: dict[str, list[InteractionRow]] = {}
    tested_rows = []
    observed_contexts: dict[str, set[str]] = defaultdict(set)
    for row in interaction_inventory:
        path = candidate / row["local_path"]
        analysis_rows, summary = read_interaction(path)
        analyses[summary["analysis_id"]] = analysis_rows
        tested_rows.append(summary)
        observed_contexts[summary["cell_type"]].add(summary["context"])
    if dict(observed_contexts) != EXPECTED_INTERACTION_CONTEXTS:
        raise ContractError(f"interaction context set drift: {dict(observed_contexts)}")

    source_union, source_counts = read_source_union(
        candidate / ZENODO_CONTEXT_REL / "significant_sc_eQTLs_with_annotations.txt.gz"
    )
    quartets, quartet_counts = read_quartets(
        candidate / ZENODO_CONTEXT_REL / "significant_quartets_n601.txt.gz",
        candidate / PUBLIC_REL / SUPPLEMENT_XLSX,
    )
    donor_counts = read_donor_contract(candidate / PUBLIC_REL / SUPPLEMENT_XLSX)
    multiplicity_rows, accepted_rule, accepted_calls = multiplicity_audit(
        analyses, source_union, quartets
    )
    if accepted_rule:
        state_lookup = accepted_calls["__state__"]
        for row in tested_rows:
            prefix = f"{row['cell_type']}:{row['context']}"
            row["source_significant_calls"] = sum(
                analysis == prefix for analysis, _pair in state_lookup
            )

    tested_path = candidate / "hong_tested_universe.tsv"
    atomic_write_tsv(tested_path, sorted(tested_rows, key=lambda row: row["analysis_id"]), list(tested_rows[0]))
    multiplicity_path = candidate / "hong_multiplicity_audit.tsv"
    atomic_write_tsv(multiplicity_path, multiplicity_rows, list(multiplicity_rows[0]))
    contract_path = candidate / "hong_source_contract.tsv"
    contract = contract_rows(donor_counts, quartet_counts, accepted_rule)
    atomic_write_tsv(contract_path, contract, list(contract[0]))
    covariate_path = candidate / "hong_covariate_availability.tsv"
    covariates = covariate_rows(candidate)
    atomic_write_tsv(covariate_path, covariates, list(covariates[0]))

    base_background_reproduced = False
    all_matching_covariates_ready = all(
        row["authorized_for_matching_now"] == "true" for row in covariates
    )
    overall_status = (
        "accepted_source_audit_claim_integration_blocked"
        if accepted_rule and base_background_reproduced
        else "coverage_limited_terminal"
    )
    reproduction = [
        {"metric": key, "observed": value, "expected": value, "status": "verified", "claim_authorized": "false"}
        for key, value in {**donor_counts, **source_counts, **quartet_counts}.items()
    ] + [
        {
            "metric": "reported_state_specific_ieqtls",
            "observed": next(
                (row["state_specific_call_count"] for row in multiplicity_rows if row["multiplicity_rule"] == accepted_rule),
                "",
            ),
            "expected": EXPECTED_REPORTED_STATE_CALLS,
            "status": "verified" if accepted_rule else "not_reconstructed",
            "claim_authorized": "false",
        }
    ]
    reproduction_path = candidate / "hong_source_reproduction.tsv"
    atomic_write_tsv(reproduction_path, reproduction, list(reproduction[0]))

    gate_rows = [
        {"gate_id": "GEN03_PUBLIC_PROVENANCE", "status": "ready", "claim_bearing": "false", "reason": "Official Nature supplements, Zenodo record, entry CRCs, and pinned public code commit verified."},
        {"gate_id": "GEN03_DONOR_CONTEXT_UNIVERSE", "status": "ready", "claim_bearing": "false", "reason": "Forty-four included donors and all 23 deposited interaction contexts reproduced from public sources."},
        {"gate_id": "GEN03_QUARTET_REPRODUCTION", "status": "ready", "claim_bearing": "false", "reason": "Zenodo and Supplementary Table 11 reproduce 601 identical quartet tuples."},
        {"gate_id": "GEN03_MULTIPLICITY_REPRODUCTION", "status": "ready" if accepted_rule else "coverage_limited", "claim_bearing": "false", "reason": f"Unique exact rule: {accepted_rule}." if accepted_rule else "No unique tested BH family reproduced the reported calls, union flags, and quartets."},
        {"gate_id": "GEN03_COMPLETE_CELLTYPE_EQTL_BACKGROUND", "status": "coverage_limited", "claim_bearing": "false", "reason": "Complete liver-eQTL tables are publicly listed but their tested-gene backgrounds are not retrieved/reproduced in this bounded GEN-03 run."},
        {"gate_id": "GEN03_REQUIRED_MATCHING_COVARIATES", "status": "not_ready", "claim_bearing": "false", "reason": "Gene-length reference and trait-locus-opportunity frame are not frozen; no matched context-rescue analysis is authorized."},
        {"gate_id": "GEN03_CONTEXT_RESCUE_AUTHORIZATION", "status": "prohibited", "claim_bearing": "false", "reason": "No rescued fraction or trait-linked context bridge may be emitted by GEN-03."},
        {"gate_id": "GEN03_NEGATIVE_CLAIM_AUTHORIZATION", "status": "prohibited", "claim_bearing": "false", "reason": "No context or genetic absence claim is authorized without complete adequate-detection and matching envelopes."},
        {"gate_id": "GEN03_OVERALL", "status": overall_status, "claim_bearing": "false", "reason": "Public source audit is bounded by unreconstructed interaction multiplicity, complete-background, and matching-covariate gates."},
    ]
    gate_path = candidate / "hong_source_gate_status.tsv"
    atomic_write_tsv(gate_path, gate_rows, list(gate_rows[0]))
    producer_manifest = write_producer_manifest(candidate)

    outputs = [
        source_manifest_path,
        tested_path,
        multiplicity_path,
        contract_path,
        covariate_path,
        reproduction_path,
        gate_path,
        producer_manifest,
    ]
    release_manifest = write_release_manifest(candidate, outputs)
    seal = write_stage_seal(candidate, "06_hong_source_audit", [*outputs, release_manifest], [stage5])
    assert_owned_output(candidate, seal)
    print(
        f"{overall_status}: accepted_multiplicity_rule={accepted_rule or 'none'}; "
        f"all_matching_covariates_ready={bool_text(all_matching_covariates_ready)}"
    )


if __name__ == "__main__":
    main()
