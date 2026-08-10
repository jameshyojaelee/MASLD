#!/usr/bin/env python3
"""Run Plan 42 source, schema, biological-unit, and access gates."""

from __future__ import annotations

import csv
import gzip
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from risk_state_common import (
    CANDIDATE_ROOT,
    PROJECT_ROOT,
    md5_file,
    read_tsv,
    require_validated_seal,
    sha256_file,
    write_tsv,
)


def parse_soft(path: Path) -> list[dict[str, str]]:
    samples: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    characteristics: list[str] = []
    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        for raw in handle:
            line = raw.rstrip("\n")
            if line.startswith("^SAMPLE = "):
                if current is not None:
                    current["characteristics"] = " | ".join(characteristics)
                    samples.append(current)
                current = {"gsm": line.split("=", 1)[1].strip()}
                characteristics = []
            elif current is not None and line.startswith("!Sample_title = "):
                current["title"] = line.split("=", 1)[1].strip()
            elif current is not None and line.startswith("!Sample_source_name_ch1 = "):
                current["source_name"] = line.split("=", 1)[1].strip()
            elif current is not None and line.startswith("!Sample_characteristics_ch1 = "):
                characteristics.append(line.split("=", 1)[1].strip())
        if current is not None:
            current["characteristics"] = " | ".join(characteristics)
            samples.append(current)
    return samples


def matrix_columns(path: Path) -> list[str]:
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8", errors="replace") as handle:
        return next(csv.reader(handle, delimiter="\t"))


def normalized_headers(sheet: openpyxl.worksheet.worksheet.Worksheet) -> tuple[int, list[str]]:
    for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        headers = ["" if value is None else str(value).strip() for value in row]
        normalized = [re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_") for value in headers]
        if any("castle" in value for value in normalized) and any("symbol" in value or value == "gene" for value in normalized):
            return index, normalized
    raise RuntimeError(f"Could not identify casTLE header in sheet {sheet.title}")


def infer_gse313544(sample: dict[str, str]) -> dict[str, str]:
    text = " | ".join([sample.get("title", ""), sample.get("characteristics", "")])
    upper = text.upper()
    genotype = ""
    for label in ("GG", "GC", "CC"):
        if re.search(rf"(?<![A-Z]){label}(?![A-Z])", upper):
            genotype = label
            break
    treatment = "palmitate" if re.search(r"PALMIT|\bPA\b", upper) else "control" if re.search(r"CONTROL|VEHICLE|\bCTRL\b|\bC\b", upper) else ""
    clone_match = re.search(r"HEPG2[_\- ]*(?:GG|GC|CC)([0-9]+)", upper)
    if clone_match is None:
        clone_match = re.search(r"CLONE\s*[:_\- ]*([A-Z0-9.]+)", upper)
    clone = f"{genotype}{clone_match.group(1)}" if clone_match and genotype else ""
    dosage = {"CC": "0", "GC": "1", "GG": "2"}.get(genotype, "")
    return {"genotype": genotype, "treatment": treatment, "clone": clone, "risk_allele_dosage": dosage}


def infer_gse158182(sample: dict[str, str]) -> dict[str, str]:
    title = sample.get("title", "")
    text = " | ".join([title, sample.get("characteristics", "")])
    upper = text.upper()
    differentiation = ""
    match = re.search(r"^D([0-9]+)[_\- ]", upper)
    if match:
        differentiation = f"D{match.group(1)}"
    line = ""
    match = re.search(r"^D[0-9]+[_\- ]([A-Z]*[0-9]{1,4})[_\- ](?:C|OA|PA)(?:\b|_)", upper)
    if match:
        line = match.group(1)
    treatment = "oleate" if re.search(r"\bOA\b|OLEIC", upper) else "palmitate" if re.search(r"\bPA\b|PALMIT", upper) else "control" if re.search(r"\bC\b|CONTROL", upper) else ""
    genotype = "I148M" if "I148M" in upper or "G342" in upper else "KO" if re.search(r"KNOCK.?OUT|\bKO\b", upper) else "WT" if re.search(r"WILD.?TYPE|\bWT\b", upper) else ""
    return {"genotype": genotype, "treatment": treatment, "line": line, "differentiation": differentiation}


def count_visium() -> tuple[int, list[str]]:
    root = PROJECT_ROOT / "Analysis/Spatial/results/starsolo_visium/HRA007511"
    sample_file = PROJECT_ROOT / "Analysis/Spatial/metadata/HRA007511_samples.txt"
    samples = [line.strip() for line in sample_file.read_text(encoding="utf-8").splitlines() if line.strip()]
    complete: list[str] = []
    for sample in samples:
        raw = root / sample / "Solo.out/Gene/raw"
        if (raw / "matrix.mtx").is_file() or (raw / "matrix.mtx.gz").is_file():
            complete.append(sample)
    return len(complete), samples


def main() -> None:
    seal = require_validated_seal()
    source = CANDIDATE_ROOT / "source"
    gates: list[dict[str, object]] = []
    schema: list[dict[str, object]] = []
    sample_manifest: list[dict[str, object]] = []

    workbook_path = source / "CLCC1_ST1/41586_2025_10064_MOESM3_ESM.xlsx"
    if workbook_path.is_file():
        identity_ok = workbook_path.stat().st_size == 4_215_920 and md5_file(workbook_path) == "c78f0e43640261bd77f26652b4fd2f93"
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        required_tokens = ["Genomic_Basal", "Batch_Basal", "Batch_MASH", "Batch_Palmitate", "Batch_Oleate"]
        sheet_ok = all(any(token in name for name in workbook.sheetnames) for token in required_tokens)
        workbook_rows = []
        schema_ok = True
        for name in workbook.sheetnames:
            if not ("Genomic_Basal" in name or "Batch_" in name):
                continue
            sheet = workbook[name]
            try:
                header_row, headers = normalized_headers(sheet)
                has_effect = any("castle_effect" in value for value in headers)
                has_score = any("castle_score" in value for value in headers)
                has_p = any(value in {"p", "p_value", "pvalue", "castle_p", "castle_p_value"} or ("castle" in value and value.endswith("_p")) for value in headers)
                passed = has_effect and has_score and has_p
            except RuntimeError:
                header_row, headers, passed = 0, [], False
            schema_ok = schema_ok and passed
            workbook_rows.append((name, sheet.max_row, header_row, passed, ",".join(headers)))
        for name, n_rows, header_row, passed, headers in workbook_rows:
            schema.append({"source_id": "CLCC1_ST1", "object": name, "n_rows": n_rows, "header_row": header_row, "passed": str(passed).lower(), "detail": headers})
        genome_rows = max((row[1] for row in workbook_rows if "Genomic_Basal" in row[0]), default=0)
        batch_rows = max((row[1] for row in workbook_rows if "Batch_Basal" in row[0]), default=0)
        passed = identity_ok and sheet_ok and schema_ok and genome_rows >= 20_000 and batch_rows >= 800
        gates.append({"source_id": "CLCC1", "gate": "official_workbook", "status": "pass" if passed else "failed_validation", "claim_scope": "direct_lipid_arm" if passed else "none", "detail": f"size_md5={identity_ok}; sheets={len(workbook.sheetnames)}; genome_rows={genome_rows}; batch_rows={batch_rows}; schema={schema_ok}"})
    else:
        gates.append({"source_id": "CLCC1", "gate": "official_workbook", "status": "skipped_source_unavailable", "claim_scope": "none", "detail": "official final workbook absent"})

    for accession, expected, infer in [
        ("GSE313544", 16, infer_gse313544),
        ("GSE158182", 28, infer_gse158182),
    ]:
        soft = source / f"{accession}_SOFT/{accession}_family.soft.gz"
        counts_name = "GSE313544_normalized_counts.txt.gz" if accession == "GSE313544" else "GSE158182_RawCounts_Subread_Genes.txt.gz"
        counts = source / f"{accession}_COUNTS/{counts_name}"
        if not soft.is_file() or not counts.is_file():
            gates.append({"source_id": accession, "gate": "sample_design", "status": "skipped_source_unavailable", "claim_scope": "none", "detail": "SOFT or expression product absent"})
            continue
        samples = parse_soft(soft)
        columns = matrix_columns(counts)
        inferred_rows = []
        for sample in samples:
            design = infer(sample)
            row = {"source_id": accession, **sample, **design, "biological_unit_status": "source_gated"}
            inferred_rows.append(row)
            sample_manifest.append(row)
        if accession == "GSE313544":
            complete = [row for row in inferred_rows if row["genotype"] and row["treatment"] and row["clone"]]
            genotype_counts = Counter(row["genotype"] for row in complete)
            design_ok = len(samples) == expected and len(complete) == expected and set(genotype_counts) == {"CC", "GC", "GG"}
            status = "pass" if design_ok else "skipped_design_unidentifiable"
            detail = f"samples={len(samples)}; fully_mapped={len(complete)}; genotype_counts={dict(genotype_counts)}; matrix_columns={len(columns)}"
        else:
            complete = [row for row in inferred_rows if row["genotype"] and row["treatment"] and row["line"] and row["differentiation"]]
            line_genotype = defaultdict(set)
            for row in complete:
                line_genotype[row["line"]].add(row["genotype"])
            design_ok = len(samples) == expected and len(complete) == expected and all(len(values) == 1 for values in line_genotype.values())
            status = "pass_source_dependent" if design_ok else "skipped_design_unidentifiable"
            detail = f"samples={len(samples)}; fully_mapped={len(complete)}; lines={len(line_genotype)}; genotype_confounded_with_line=true; matrix_columns={len(columns)}"
        gates.append({"source_id": accession, "gate": "sample_design", "status": status, "claim_scope": "stress_response_only" if design_ok else "none", "detail": detail})

    n_visium, visium_samples = count_visium()
    gates.append({"source_id": "HMSMA_VISIUM", "gate": "35_matrix_completion", "status": "pass" if n_visium == 35 else "pending_repair", "claim_scope": "transcriptomic_only" if n_visium == 35 else "none", "detail": f"complete={n_visium}/35; expected_samples={len(visium_samples)}"})

    clinical_path = PROJECT_ROOT / "data/HRA007511/metadata/hmsma_st_sample_metadata.csv"
    clinical_ids: set[str] = set()
    if clinical_path.is_file():
        with clinical_path.open(newline="", encoding="utf-8") as handle:
            clinical_ids = {row["sample"] for row in csv.DictReader(handle) if row.get("sample")}
    joined_ids = set(visium_samples) & clinical_ids
    clinical_key_status = "pass" if len(joined_ids) == 35 else "skipped_no_authoritative_sample_key"
    gates.append({
        "source_id": "HMSMA_VISIUM",
        "gate": "GSA_to_clinical_sample_key",
        "status": clinical_key_status,
        "claim_scope": "histology_and_disease_contrasts" if clinical_key_status == "pass" else "unsupervised_only",
        "detail": f"GSA_ids={len(visium_samples)}; clinical_ids={len(clinical_ids)}; direct_authoritative_joins={len(joined_ids)}; HRA_xx identifiers cannot be inferred as CTRL/MASL/MASH identifiers",
    })

    portal = source / "HMSMA_PORTAL_DOWNLOAD_PAGE/hmsma_download.html"
    access_js = source / "HMSMA_PORTAL_ACCESS_JS/hmsma.popup.js"
    public_listing = portal.is_file() and "MALDI_allsamples.h5ad" in portal.read_text(encoding="utf-8", errors="replace")
    controlled_portal = access_js.is_file() and "database is under control" in access_js.read_text(encoding="utf-8", errors="replace")
    local_maldi = source / "HMSMA_MALDI/MALDI_allsamples.h5ad"
    if local_maldi.is_file():
        maldi_status = "pending_schema_audit"
        detail = f"local_public_product_present=true; bytes={local_maldi.stat().st_size}; authoritative_registration_unverified"
    else:
        maldi_status = "skipped_no_public_registered_maldi"
        detail = f"portal_lists_38G_h5ad={public_listing}; portal_explicitly_controlled={controlled_portal}; OMIX009098_controlled_access=true; no no-permission object or authoritative registration recovered"
    gates.append({"source_id": "HMSMA_MALDI", "gate": "public_registered_product", "status": maldi_status, "claim_scope": "none", "detail": detail})

    write_tsv(CANDIDATE_ROOT / "source_gate_status.tsv", gates, ["source_id", "gate", "status", "claim_scope", "detail"])
    write_tsv(CANDIDATE_ROOT / "source_schema_audit.tsv", schema, ["source_id", "object", "n_rows", "header_row", "passed", "detail"])
    sample_fields = sorted({key for row in sample_manifest for key in row}) if sample_manifest else ["source_id"]
    write_tsv(CANDIDATE_ROOT / "sample_manifest.tsv", sample_manifest, sample_fields)
    result = {"candidate_id": seal["candidate_id"], "gates": gates, "source_manifest_sha256": sha256_file(CANDIDATE_ROOT / "source_manifest.tsv") if (CANDIDATE_ROOT / "source_manifest.tsv").is_file() else ""}
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
